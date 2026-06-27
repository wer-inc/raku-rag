"""Datasource onboarding preview: mapping suggestions, normalization, and validation.

This layer is deliberately before ingestion. It samples connector output and explains how customer
columns would map into a canonical vocabulary without indexing anything. The vocabulary is
profile-driven so FAQ, manuals, and manufacturing tables can be onboarded by configuration instead
of per-customer code.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date
from typing import Any, Mapping

from raku_rag.services.datasource_sync import SyncDocument

CANONICAL_MANUFACTURING_FIELDS = (
    "equipment_id",
    "factory_id",
    "line_id",
    "process_id",
    "alarm_code",
    "defect_type",
    "part_no",
    "document_kind",
    "approval_status",
    "effective_date",
    "acl_tags",
)

CANONICAL_CONTENT_FIELDS = (
    "question",
    "answer",
    "category",
    "document_id",
    "document_title",
    "document_version",
    "section",
    "heading",
    "body",
    "product_id",
    "product_name",
    "owner_department",
    "published_at",
    "updated_at",
)

CANONICAL_ALL_FIELDS = tuple(
    dict.fromkeys((*CANONICAL_MANUFACTURING_FIELDS, *CANONICAL_CONTENT_FIELDS))
)

PROFILE_LABELS = {
    "auto": "自動判定",
    "manufacturing": "設備・工程データ",
    "faq": "FAQ",
    "manual": "マニュアル・手順書",
    "generic": "汎用文書",
}

CANONICAL_FIELDS_BY_PROFILE = {
    "manufacturing": CANONICAL_MANUFACTURING_FIELDS,
    "faq": (
        "question",
        "answer",
        "category",
        "product_id",
        "product_name",
        "equipment_id",
        "factory_id",
        "owner_department",
        "published_at",
        "updated_at",
        "document_kind",
        "approval_status",
        "effective_date",
        "acl_tags",
    ),
    "manual": (
        "document_id",
        "document_title",
        "document_version",
        "section",
        "heading",
        "body",
        "category",
        "product_id",
        "product_name",
        "equipment_id",
        "factory_id",
        "line_id",
        "owner_department",
        "published_at",
        "updated_at",
        "document_kind",
        "approval_status",
        "effective_date",
        "acl_tags",
    ),
    "generic": (
        "document_id",
        "document_title",
        "category",
        "body",
        "owner_department",
        "published_at",
        "updated_at",
        "document_kind",
        "approval_status",
        "effective_date",
        "acl_tags",
    ),
}

CANONICAL_SECTIONS_BY_PROFILE = {
    "manufacturing": (
        ("設備・工程", ("equipment_id", "factory_id", "line_id", "process_id", "part_no")),
        ("トラブル", ("alarm_code", "defect_type")),
        ("ガバナンス", ("document_kind", "approval_status", "effective_date", "acl_tags")),
    ),
    "faq": (
        ("FAQ", ("question", "answer", "category")),
        ("対象", ("product_id", "product_name", "equipment_id", "factory_id")),
        ("管理", ("owner_department", "published_at", "updated_at")),
        ("ガバナンス", ("document_kind", "approval_status", "effective_date", "acl_tags")),
    ),
    "manual": (
        (
            "文書",
            (
                "document_id",
                "document_title",
                "document_version",
                "section",
                "heading",
                "body",
            ),
        ),
        ("対象", ("category", "product_id", "product_name", "equipment_id", "factory_id", "line_id")),
        ("管理", ("owner_department", "published_at", "updated_at")),
        ("ガバナンス", ("document_kind", "approval_status", "effective_date", "acl_tags")),
    ),
    "generic": (
        ("文書", ("document_id", "document_title", "category", "body")),
        ("管理", ("owner_department", "published_at", "updated_at")),
        ("ガバナンス", ("document_kind", "approval_status", "effective_date", "acl_tags")),
    ),
}

PROFILE_DEFAULT_REQUIRED_FIELDS = {
    "manufacturing": ("equipment_id",),
    "faq": ("question", "answer"),
    "manual": (),
    "generic": (),
}

DEFAULT_REQUIRED_FIELDS = PROFILE_DEFAULT_REQUIRED_FIELDS["manufacturing"]

_FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "equipment_id": (
        "equipment_id",
        "equipment id",
        "equipment",
        "machine id",
        "machine code",
        "machinecode",
        "asset id",
        "assetid",
        "設備番号",
        "設備id",
        "設備ID",
        "設備コード",
        "機番",
    ),
    "factory_id": ("factory_id", "factory", "plant", "site", "工場", "工場id", "拠点"),
    "line_id": ("line_id", "line", "production line", "ライン", "生産ライン"),
    "process_id": ("process_id", "process", "工程", "工程id", "プロセス"),
    "alarm_code": ("alarm_code", "alarm", "alarm code", "alarmcode", "アラーム", "アラームコード"),
    "defect_type": ("defect_type", "defect", "defect category", "不具合", "不具合分類", "欠陥種別"),
    "part_no": ("part_no", "part number", "part_number", "partno", "品番", "部品番号", "型番"),
    "document_kind": ("document_kind", "document type", "document_type", "文書種別", "資料種別"),
    "approval_status": ("approval_status", "approval", "status", "承認状態", "承認ステータス"),
    "effective_date": ("effective_date", "effective date", "valid from", "有効日", "適用日"),
    "acl_tags": ("acl_tags", "acl", "access", "permission", "権限", "アクセス範囲"),
    "question": (
        "question",
        "q",
        "faq_question",
        "question_text",
        "問い合わせ",
        "問い合わせ内容",
        "質問",
        "質問文",
        "問合せ",
        "件名",
    ),
    "answer": (
        "answer",
        "a",
        "faq_answer",
        "answer_text",
        "response",
        "reply",
        "回答",
        "回答文",
        "対応内容",
        "最終回答",
    ),
    "category": ("category", "カテゴリ", "カテゴリー", "分類", "区分", "タグ"),
    "document_id": (
        "document_id",
        "doc_id",
        "docid",
        "document code",
        "document_code",
        "文書id",
        "文書ID",
        "文書番号",
        "資料番号",
    ),
    "document_title": (
        "document_title",
        "title",
        "manual_title",
        "document name",
        "document_name",
        "タイトル",
        "文書名",
        "資料名",
        "マニュアル名",
    ),
    "document_version": (
        "document_version",
        "version",
        "revision",
        "rev",
        "版",
        "版数",
        "改訂",
        "改訂番号",
        "バージョン",
    ),
    "section": ("section", "chapter", "章", "節", "セクション", "章節"),
    "heading": ("heading", "header", "見出し", "小見出し", "項目名"),
    "body": (
        "body",
        "content",
        "text",
        "description",
        "detail",
        "details",
        "本文",
        "内容",
        "詳細",
        "説明",
        "手順",
    ),
    "product_id": (
        "product_id",
        "product id",
        "product_code",
        "product code",
        "sku",
        "製品id",
        "製品ID",
        "製品コード",
        "商品id",
        "商品ID",
    ),
    "product_name": ("product_name", "product name", "product", "製品名", "商品名", "製品"),
    "owner_department": (
        "owner_department",
        "department",
        "owner",
        "担当部署",
        "管理部署",
        "部門",
        "部署",
    ),
    "published_at": ("published_at", "published", "publish_date", "公開日", "発行日", "掲載日"),
    "updated_at": ("updated_at", "updated", "modified_at", "updated date", "更新日", "更新日時"),
}


def _normalize_key(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).casefold()
    return re.sub(r"[\s_\-./:：()（）\[\]【】]+", "", text)


_ALIAS_TO_FIELD = {
    _normalize_key(alias): field
    for field, aliases in _FIELD_ALIASES.items()
    for alias in aliases
}

_APPROVAL_STATUS_ALIASES = {
    "approved": "approved",
    "approve": "approved",
    "承認済": "approved",
    "承認済み": "approved",
    "draft": "draft",
    "下書き": "draft",
    "ドラフト": "draft",
    "pending": "pending_review",
    "pendingreview": "pending_review",
    "pending_review": "pending_review",
    "review": "pending_review",
    "要確認": "pending_review",
    "要レビュー": "pending_review",
    "obsolete": "obsolete",
    "廃止": "obsolete",
    "旧版": "obsolete",
}

_DOCUMENT_KIND_ALIASES = {
    "workinstruction": "work_instruction",
    "work_instruction": "work_instruction",
    "手順書": "work_instruction",
    "作業手順": "work_instruction",
    "inspection": "inspection",
    "検査": "inspection",
    "qualityreport": "quality_report",
    "quality_report": "quality_report",
    "品質報告": "quality_report",
    "troublereport": "trouble_report",
    "trouble_report": "trouble_report",
    "トラブル報告": "trouble_report",
    "minutes": "minutes",
    "議事録": "minutes",
    "ledger": "ledger",
    "台帳": "ledger",
    "drawing": "drawing",
    "図面": "drawing",
    "training": "training",
    "教育": "training",
}


def build_datasource_preview(
    *,
    source_id: str,
    datasource: Mapping[str, object],
    documents: list[SyncDocument],
    body: Mapping[str, object] | None = None,
) -> dict[str, object]:
    """Build a serializable onboarding preview from already-fetched sample documents."""

    body = body or {}
    profile = _mapping_profile(datasource, body)
    max_rows = _int_value(body, profile, ("sample_rows", "max_sample_rows"), default=5)
    max_rows = max(1, min(max_rows, 20))

    tables: list[dict[str, object]] = []
    document_summaries: list[dict[str, object]] = []
    detected_columns: list[str] = []
    for document in documents:
        parsed_tables = _tables_from_document(document)
        document_summaries.append(
            {
                "document_id": document.document_id,
                "document_ref": document.document_ref,
                "content_type": document.content_type,
                "kind": "table" if parsed_tables else "text",
                "sample_row_count": sum(len(table["rows"]) for table in parsed_tables),
                "text_preview": _text_preview(document.raw, document.content_type)
                if not parsed_tables
                else "",
            }
        )
        for table in parsed_tables:
            tables.append(table)
            for column in table["columns"]:
                if column not in detected_columns:
                    detected_columns.append(str(column))

    profile_type = _resolve_profile_type(profile, detected_columns, document_summaries)
    canonical_fields = _canonical_fields(profile_type)
    explicit_mapping = _explicit_mapping(profile, canonical_fields)
    suggested_mapping, mapping_confidence = _suggest_mapping(
        detected_columns,
        explicit_mapping,
        canonical_fields,
    )
    defaults = _defaults(profile, canonical_fields)
    required_fields = _required_fields(profile, canonical_fields, profile_type)

    sample_rows: list[dict[str, object]] = []
    validation_errors: list[str] = []
    validation_warnings: list[str] = []
    valid_count = 0
    needs_review_count = 0
    for table in tables:
        for raw_row in table["rows"][:max_rows]:
            normalized = _normalize_row(raw_row["values"], suggested_mapping, defaults)
            row_errors, row_warnings = _validate_row(normalized, required_fields)
            if row_errors:
                needs_review_count += 1
            else:
                valid_count += 1
            validation_errors.extend(
                f"row {raw_row['row_number']}: {error}" for error in row_errors
            )
            validation_warnings.extend(
                f"row {raw_row['row_number']}: {warning}" for warning in row_warnings
            )
            sample_rows.append(
                {
                    "document_id": table["document_id"],
                    "sheet_name": table["sheet_name"],
                    "row_number": raw_row["row_number"],
                    "raw": raw_row["values"],
                    "normalized": normalized,
                    "validation": {
                        "status": "needs_review" if row_errors else "valid",
                        "errors": row_errors,
                        "warnings": row_warnings,
                    },
                }
            )
            if len(sample_rows) >= max_rows:
                break
        if len(sample_rows) >= max_rows:
            break

    unmapped = [column for column in detected_columns if column not in suggested_mapping]
    if detected_columns and not suggested_mapping:
        validation_warnings.append("no columns mapped to canonical fields")
    elif unmapped:
        validation_warnings.append(
            "unmapped columns: " + ", ".join(unmapped[:8])
        )

    return {
        "source_id": source_id,
        "profile_type": profile_type,
        "profile_label": PROFILE_LABELS[profile_type],
        "profile_options": _profile_options(),
        "canonical_sections": _canonical_sections(profile_type),
        "document_count": len(documents),
        "documents": document_summaries,
        "canonical_fields": list(canonical_fields),
        "detected_columns": detected_columns,
        "explicit_mapping": explicit_mapping,
        "suggested_mapping": suggested_mapping,
        "mapping_confidence": mapping_confidence,
        "defaults": defaults,
        "required_fields": list(required_fields),
        "sample_rows": sample_rows,
        "validation": {
            "valid_count": valid_count,
            "needs_review_count": needs_review_count,
            "error_count": len(validation_errors),
            "warning_count": len(validation_warnings),
            "errors": validation_errors[:25],
            "warnings": list(dict.fromkeys(validation_warnings))[:25],
        },
    }


def _mapping_profile(
    datasource: Mapping[str, object], body: Mapping[str, object]
) -> dict[str, object]:
    config = _mapping(datasource.get("config"))
    profile: dict[str, object] = {}
    for source in (config, _mapping(config.get("mapping_profile")), body, _mapping(body.get("mapping_profile"))):
        for key in (
            "profile_type",
            "data_profile",
            "field_mapping",
            "mapping",
            "defaults",
            "required_fields",
            "sample_rows",
        ):
            if key in source:
                if key in {"field_mapping", "mapping", "defaults"}:
                    merged = dict(profile.get(key) or {})
                    merged.update(_mapping(source.get(key)))
                    profile[key] = merged
                else:
                    profile[key] = source[key]
    return profile


def _resolve_profile_type(
    profile: Mapping[str, object],
    detected_columns: list[str],
    document_summaries: list[dict[str, object]],
) -> str:
    requested = _normalize_profile_type(
        profile.get("profile_type") or profile.get("data_profile") or "auto"
    )
    if requested != "auto":
        return requested

    matched_fields = _matched_fields(detected_columns)
    if {"question", "answer"}.issubset(matched_fields):
        return "faq"
    if matched_fields.intersection({"question", "answer"}) and not matched_fields.intersection(
        CANONICAL_MANUFACTURING_FIELDS
    ):
        return "faq"
    if matched_fields.intersection(
        {"document_id", "document_title", "document_version", "section", "heading", "body"}
    ):
        return "manual"
    if matched_fields.intersection(CANONICAL_MANUFACTURING_FIELDS):
        return "manufacturing"
    if any(summary.get("kind") == "text" for summary in document_summaries):
        return "manual"
    return "generic"


def _normalize_profile_type(value: object) -> str:
    normalized = _normalize_key(value)
    aliases = {
        "": "auto",
        "auto": "auto",
        "detect": "auto",
        "自動": "auto",
        "自動判定": "auto",
        "manufacturing": "manufacturing",
        "mfg": "manufacturing",
        "factory": "manufacturing",
        "equipment": "manufacturing",
        "設備": "manufacturing",
        "製造": "manufacturing",
        "faq": "faq",
        "qa": "faq",
        "qanda": "faq",
        "質問回答": "faq",
        "manual": "manual",
        "manuals": "manual",
        "document": "manual",
        "docs": "manual",
        "手順書": "manual",
        "マニュアル": "manual",
        "generic": "generic",
        "general": "generic",
        "汎用": "generic",
    }
    return aliases.get(normalized, "auto")


def _matched_fields(columns: list[str]) -> set[str]:
    matched: set[str] = set()
    for column in columns:
        normalized = _normalize_key(column)
        direct = _ALIAS_TO_FIELD.get(normalized)
        if direct:
            matched.add(direct)
            continue
        for alias, canonical in _ALIAS_TO_FIELD.items():
            if alias and alias in normalized:
                matched.add(canonical)
    return matched


def _canonical_fields(profile_type: str) -> tuple[str, ...]:
    return tuple(CANONICAL_FIELDS_BY_PROFILE.get(profile_type, CANONICAL_FIELDS_BY_PROFILE["generic"]))


def _profile_options() -> list[dict[str, str]]:
    return [
        {"profile_type": key, "label": label}
        for key, label in PROFILE_LABELS.items()
    ]


def _canonical_sections(profile_type: str) -> list[dict[str, object]]:
    fields = set(_canonical_fields(profile_type))
    sections = CANONICAL_SECTIONS_BY_PROFILE.get(profile_type, CANONICAL_SECTIONS_BY_PROFILE["generic"])
    return [
        {"label": label, "fields": [field for field in section_fields if field in fields]}
        for label, section_fields in sections
    ]


def _explicit_mapping(
    profile: Mapping[str, object], canonical_fields: tuple[str, ...]
) -> dict[str, str]:
    raw_mapping = {}
    raw_mapping.update(_mapping(profile.get("mapping")))
    raw_mapping.update(_mapping(profile.get("field_mapping")))
    allowed = set(canonical_fields)
    explicit: dict[str, str] = {}
    for source_column, canonical_field in raw_mapping.items():
        field = str(canonical_field)
        if field in allowed:
            explicit[str(source_column)] = field
    return explicit


def _defaults(
    profile: Mapping[str, object], canonical_fields: tuple[str, ...]
) -> dict[str, object]:
    defaults: dict[str, object] = {}
    allowed = set(canonical_fields)
    for key, value in _mapping(profile.get("defaults")).items():
        if key in allowed and value not in (None, ""):
            defaults[str(key)] = _normalize_value(str(key), value)
    return defaults


def _required_fields(
    profile: Mapping[str, object],
    canonical_fields: tuple[str, ...],
    profile_type: str,
) -> tuple[str, ...]:
    raw = profile.get("required_fields")
    if isinstance(raw, str):
        fields = [part.strip() for part in raw.split(",")]
    elif isinstance(raw, (list, tuple)):
        fields = [str(part).strip() for part in raw]
    else:
        fields = list(PROFILE_DEFAULT_REQUIRED_FIELDS.get(profile_type, ()))
    allowed = set(canonical_fields)
    return tuple(field for field in fields if field in allowed)


def _suggest_mapping(
    columns: list[str],
    explicit: Mapping[str, str],
    canonical_fields: tuple[str, ...],
) -> tuple[dict[str, str], dict[str, float]]:
    mapping: dict[str, str] = {}
    confidence: dict[str, float] = {}
    allowed = set(canonical_fields)
    for column in columns:
        if column in explicit:
            mapping[column] = explicit[column]
            confidence[column] = 1.0
            continue
        normalized = _normalize_key(column)
        if normalized in _ALIAS_TO_FIELD and _ALIAS_TO_FIELD[normalized] in allowed:
            mapping[column] = _ALIAS_TO_FIELD[normalized]
            confidence[column] = 0.92
            continue
        for alias, canonical in _ALIAS_TO_FIELD.items():
            if alias and alias in normalized and canonical in allowed:
                mapping[column] = canonical
                confidence[column] = 0.72
                break
    return mapping, confidence


def _normalize_row(
    row: Mapping[str, object], mapping: Mapping[str, str], defaults: Mapping[str, object]
) -> dict[str, object]:
    normalized: dict[str, object] = {}
    for source_column, raw_value in row.items():
        canonical = mapping.get(source_column)
        if canonical:
            normalized[canonical] = _normalize_value(canonical, raw_value)
    for key, value in defaults.items():
        if normalized.get(key) in (None, ""):
            normalized[key] = value
    return normalized


def _normalize_value(field: str, value: object) -> object:
    text = _clean_cell(value)
    if field == "approval_status":
        return _APPROVAL_STATUS_ALIASES.get(_normalize_key(text), text)
    if field == "document_kind":
        return _DOCUMENT_KIND_ALIASES.get(_normalize_key(text), text)
    if field in {"effective_date", "published_at", "updated_at"}:
        return text.replace("/", "-")
    if field == "acl_tags":
        if isinstance(value, (list, tuple)):
            return [str(item).strip() for item in value if str(item).strip()]
        return [part.strip() for part in re.split(r"[,/;、\s]+", text) if part.strip()]
    return text


def _validate_row(
    normalized: Mapping[str, object], required_fields: tuple[str, ...]
) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    for field in required_fields:
        if normalized.get(field) in (None, "", []):
            errors.append(f"missing required field: {field}")
    status = normalized.get("approval_status")
    if status and status not in {"draft", "pending_review", "approved", "obsolete"}:
        warnings.append(f"unknown approval_status: {status}")
    for date_field in ("effective_date", "published_at", "updated_at"):
        value = normalized.get(date_field)
        if not value:
            continue
        try:
            date.fromisoformat(str(value))
        except ValueError:
            warnings.append(f"invalid {date_field}: {value}")
    return errors, warnings


def _tables_from_document(document: SyncDocument) -> list[dict[str, object]]:
    if document.content_type in {"text/csv", "application/csv"}:
        return [_csv_table(document)]
    if document.content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return _xlsx_tables(document)
    return []


def _csv_table(document: SyncDocument) -> dict[str, object]:
    text = document.raw.decode("utf-8-sig", errors="replace")
    rows = [list(row) for row in csv.reader(io.StringIO(text))]
    return _table_from_rows(document, "sheet1", rows)


def _xlsx_tables(document: SyncDocument) -> list[dict[str, object]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:  # pragma: no cover - optional dependency depends on install profile
        raise RuntimeError("openpyxl is required for XLSX onboarding preview") from exc

    wb = load_workbook(io.BytesIO(document.raw), read_only=True, data_only=True)
    try:
        tables = []
        for ws in wb.worksheets:
            rows = [
                ["" if value is None else str(value) for value in row]
                for row in ws.iter_rows(values_only=True)
            ]
            tables.append(_table_from_rows(document, ws.title, rows))
        return tables
    finally:
        wb.close()


def _table_from_rows(
    document: SyncDocument, sheet_name: str, rows: list[list[object]]
) -> dict[str, object]:
    header = [_clean_cell(value) for value in (rows[0] if rows else [])]
    columns = [name or f"column_{idx}" for idx, name in enumerate(header, start=1)]
    data_rows: list[dict[str, object]] = []
    for row_number, row in enumerate(rows[1:], start=2):
        values = {}
        for index, value in enumerate(row, start=1):
            column = columns[index - 1] if index - 1 < len(columns) else f"column_{index}"
            values[column] = _clean_cell(value)
        if any(value not in ("", None) for value in values.values()):
            data_rows.append({"row_number": row_number, "values": values})
    return {
        "document_id": document.document_id,
        "document_ref": document.document_ref,
        "sheet_name": sheet_name,
        "columns": columns,
        "rows": data_rows,
    }


def _text_preview(raw: bytes, content_type: str) -> str:
    text = raw.decode("utf-8", errors="replace")
    if content_type == "text/html":
        text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text)).strip()
    return text[:240]


def _mapping(value: object) -> Mapping[str, object]:
    return value if isinstance(value, dict) else {}


def _int_value(
    body: Mapping[str, object],
    profile: Mapping[str, object],
    keys: tuple[str, ...],
    *,
    default: int,
) -> int:
    for source in (body, profile):
        for key in keys:
            try:
                value = source.get(key)
            except AttributeError:
                continue
            if value not in (None, ""):
                try:
                    return int(value)
                except (TypeError, ValueError):
                    return default
    return default


def _clean_cell(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    return re.sub(r"[ \t]+", " ", text).strip()
