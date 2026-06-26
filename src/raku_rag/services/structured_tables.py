"""Structured table manifests and deterministic table-query tool.

CSV/XLSX data is indexed as searchable cell chunks by ``SpreadsheetParser``. This module keeps the
separate structured view needed for exact table operations so RAG does not pretend vector search is a
database for aggregation, ranking, latest-value, or numeric-filter questions.
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from datetime import datetime
from statistics import median
from typing import Any, Mapping

from raku_rag.core.errors import AnswerStatus
from raku_rag.domain.models import Answer, Citation, Freshness, IdentityClaims
from raku_rag.providers.parsers import (
    CSV_CONTENT_TYPES,
    XLSX_CONTENT_TYPE,
    cell_anchor,
    parse_cell_anchor,
)

STRUCTURED_TABLE_MANIFEST_VERSION = "structured-table-manifest-v1"
STRUCTURED_TABLE_MANIFESTS_KEY = "structured_table_manifests"
STRUCTURED_TABLE_MANIFEST_VERSION_KEY = "structured_table_manifest_version"
STRUCTURED_TABLE_COUNT_KEY = "structured_table_count"

_WORD = re.compile(r"[A-Za-z0-9_一-龯ぁ-んァ-ン]+")
_TOP_N = re.compile(r"\btop\s+(\d+)\b", re.I)
_NUMBER = re.compile(r"[-+]?\d+(?:,\d{3})*(?:\.\d+)?")
_YEAR_MONTH = re.compile(r"\b(\d{4})[-/](\d{1,2})\b")
_YEAR = re.compile(r"\b(20\d{2}|19\d{2})\b")
_MONTHS = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}


@dataclass(frozen=True)
class _CellRef:
    document_id: str
    source_id: str
    version: int
    indexed_at: str
    sheet_name: str
    row_id: str
    cell_range: str


@dataclass(frozen=True)
class _StructuredRow:
    document_id: str
    source_id: str
    version: int
    indexed_at: str
    sheet_name: str
    row_id: str
    row_number: int
    values: Mapping[str, str]
    cells: Mapping[str, str]


def extract_structured_table_manifests(raw: bytes, content_type: str) -> list[dict[str, Any]]:
    """Return serializable table manifests for CSV/XLSX inputs; non-table inputs return ``[]``."""

    if content_type == XLSX_CONTENT_TYPE:
        sheets = _read_xlsx(raw)
    elif content_type in CSV_CONTENT_TYPES:
        sheets = _read_csv(raw)
    else:
        return []
    manifests: list[dict[str, Any]] = []
    for sheet_name, rows in sheets:
        if not rows:
            continue
        header = [_clean_cell(value) for value in rows[0]]
        columns = [
            {"name": name or f"column_{index}", "index": index}
            for index, name in enumerate(header, start=1)
        ]
        data_rows: list[dict[str, Any]] = []
        for row_number, row in enumerate(rows[1:], start=2):
            values: dict[str, str] = {}
            cells: dict[str, str] = {}
            for col_index, raw_value in enumerate(row, start=1):
                value = _clean_cell(raw_value)
                if not value:
                    continue
                column = header[col_index - 1] if col_index - 1 < len(header) else ""
                column = column or f"column_{col_index}"
                values[column] = value
                cells[column] = cell_anchor(sheet_name, row_number, col_index)
            if values:
                data_rows.append(
                    {
                        "row_id": str(row_number),
                        "row_number": row_number,
                        "values": values,
                        "cells": cells,
                    }
                )
        if data_rows:
            manifests.append(
                {
                    "version": STRUCTURED_TABLE_MANIFEST_VERSION,
                    "table_id": sheet_name,
                    "sheet_name": sheet_name,
                    "header_row": 1,
                    "columns": columns,
                    "rows": data_rows,
                }
            )
    return manifests


def redact_table_manifest_values(
    manifests: Iterable[Mapping[str, Any]], redact: Callable[[str], str]
) -> list[dict[str, Any]]:
    """Apply the same pre-index redaction policy to stored cell values."""

    out: list[dict[str, Any]] = []
    for manifest in manifests:
        copied = dict(manifest)
        rows = []
        for row in manifest.get("rows") or []:
            if not isinstance(row, Mapping):
                continue
            values = {
                str(key): redact(str(value)) for key, value in dict(row.get("values") or {}).items()
            }
            copied_row = dict(row)
            copied_row["values"] = values
            rows.append(copied_row)
        copied["rows"] = rows
        out.append(copied)
    return out


def cell_metadata_for_text(text: str, manifests: Iterable[Mapping[str, Any]]) -> dict[str, object]:
    anchor = parse_cell_anchor(text)
    if anchor is None:
        return {}
    sheet, row, col = anchor
    cell = f"{sheet}!R{row}C{col}"
    column = _column_name_for_cell(cell, manifests)
    return {
        "structured_kind": "spreadsheet_cell",
        "sheet_name": sheet,
        "row_id": str(row),
        "column_index": col,
        "column_name": column,
        "cell_range": cell,
    }


class TableManifestStructuredTool:
    """Deterministic structured-query tool over tenant-visible table manifests."""

    def __init__(self, registry, acl_policy, *, max_result_rows: int = 10) -> None:
        self._registry = registry
        self._acl = acl_policy
        self._max_result_rows = max(1, int(max_result_rows))

    def answer(
        self,
        *,
        principal: IdentityClaims,
        query: str,
        collection_id: str | None = None,
        reason: str = "",
        correlation_id: str = "",
    ) -> Answer:
        rows = list(self._visible_rows(principal, collection_id=collection_id))
        if not rows:
            return Answer(
                status=AnswerStatus.INSUFFICIENT_EVIDENCE.value,
                text=None,
                citations=(),
                used_chunks=(),
                correlation_id=correlation_id,
                route="structured_tool",
            )
        result = self._run(query, reason, rows)
        if not result:
            return Answer(
                status=AnswerStatus.INSUFFICIENT_EVIDENCE.value,
                text=None,
                citations=(),
                used_chunks=(),
                correlation_id=correlation_id,
                route="structured_tool",
            )
        text, refs = result
        citations = tuple(_citation(ref) for ref in refs[: self._max_result_rows])
        used = tuple(f"{ref.document_id}:table:{ref.sheet_name}:R{ref.row_id}" for ref in refs)
        freshness = tuple(
            Freshness(indexed_at=ref.indexed_at, document_version=ref.version)
            for ref in refs[: self._max_result_rows]
        )
        return Answer(
            status=AnswerStatus.OK.value,
            text=text,
            confidence=1.0,
            citations=citations,
            used_chunks=used,
            freshness=freshness,
            correlation_id=correlation_id,
            route="structured_tool",
        )

    def _run(
        self, query: str, reason: str, rows: list[_StructuredRow]
    ) -> tuple[str, list[_CellRef]] | None:
        if reason == "aggregation":
            return _aggregate(query, rows)
        if reason == "ranking":
            return _rank(query, rows, self._max_result_rows)
        if reason == "latest_value":
            return _latest(query, rows)
        if reason == "numeric_comparison":
            return _numeric_filter(query, rows, self._max_result_rows)
        if reason == "period_filter":
            return _period_filter(query, rows, self._max_result_rows)
        return None

    def _visible_rows(
        self, principal: IdentityClaims, *, collection_id: str | None = None
    ) -> Iterable[_StructuredRow]:
        for doc in _tenant_documents(self._registry, principal.tenant_id):
            if doc.tombstone:
                continue
            if collection_id and doc.collection_id != collection_id:
                continue
            if not self._acl.can_read_document(principal, doc):
                continue
            for manifest in doc.metadata.get(STRUCTURED_TABLE_MANIFESTS_KEY) or []:
                if not isinstance(manifest, Mapping):
                    continue
                sheet = str(manifest.get("sheet_name") or manifest.get("table_id") or "sheet1")
                for row in manifest.get("rows") or []:
                    if not isinstance(row, Mapping):
                        continue
                    values = {str(k): str(v) for k, v in dict(row.get("values") or {}).items()}
                    cells = {str(k): str(v) for k, v in dict(row.get("cells") or {}).items()}
                    if not values:
                        continue
                    yield _StructuredRow(
                        document_id=doc.document_id,
                        source_id=doc.source_id,
                        version=doc.version,
                        indexed_at=doc.indexed_at,
                        sheet_name=sheet,
                        row_id=str(row.get("row_id") or row.get("row_number") or ""),
                        row_number=int(row.get("row_number") or 0),
                        values=values,
                        cells=cells,
                    )


def _read_csv(raw: bytes) -> list[tuple[str, list[list[str]]]]:
    text = raw.decode("utf-8-sig", errors="replace")
    return [("sheet1", [list(row) for row in csv.reader(io.StringIO(text))])]


def _read_xlsx(raw: bytes) -> list[tuple[str, list[list[str]]]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sheets = []
        for ws in wb.worksheets:
            rows = [
                ["" if value is None else str(value) for value in row]
                for row in ws.iter_rows(values_only=True)
            ]
            sheets.append((ws.title, rows))
        return sheets
    finally:
        wb.close()


def _clean_cell(value: object) -> str:
    return re.sub(r"[ \t]+", " ", str(value or "")).strip()


def _tenant_documents(registry, tenant_id: str):
    if hasattr(registry, "documents_for_tenant"):
        return registry.documents_for_tenant(tenant_id)
    if hasattr(registry, "list_documents"):
        return tuple(registry.list_documents(tenant_id))
    return ()


def _column_name_for_cell(cell: str, manifests: Iterable[Mapping[str, Any]]) -> str:
    for manifest in manifests:
        for row in manifest.get("rows") or []:
            cells = dict(row.get("cells") or {}) if isinstance(row, Mapping) else {}
            for column, candidate in cells.items():
                if candidate == cell:
                    return str(column)
    return ""


def _citation(ref: _CellRef) -> Citation:
    return Citation(
        kind="spreadsheet",
        document_id=ref.document_id,
        source_id=ref.source_id,
        version=ref.version,
        retrieval_score=1.0,
        chunk_id=f"{ref.document_id}:table:{ref.sheet_name}:R{ref.row_id}",
        sheet_name=ref.sheet_name,
        cell_range=ref.cell_range,
        row_id=ref.row_id,
    )


def _aggregate(query: str, rows: list[_StructuredRow]) -> tuple[str, list[_CellRef]] | None:
    op = _aggregate_op(query)
    value_col = _best_numeric_column(query, rows)
    if not value_col and op != "count":
        return None
    group_col = _group_column(query, rows, exclude=value_col)
    if group_col:
        grouped: dict[str, list[tuple[_StructuredRow, float]]] = {}
        for row in rows:
            key = row.values.get(group_col, "")
            value = 1.0 if op == "count" else _numeric(row.values.get(value_col, ""))
            if key and value is not None:
                grouped.setdefault(key, []).append((row, value))
        if not grouped:
            return None
        parts = []
        refs: list[_CellRef] = []
        for key, row_values in grouped.items():
            values = [value for _row, value in row_values]
            parts.append(f"{key}={_format_number(_apply_aggregate(op, values))}")
            refs.extend(_refs_for_rows([row for row, _value in row_values], value_col or group_col))
        label = value_col if op != "count" else "rows"
        return f"{_op_label(op)} {label} by {group_col}: " + "; ".join(parts), refs

    values_and_rows = []
    for row in rows:
        value = 1.0 if op == "count" else _numeric(row.values.get(value_col, ""))
        if value is not None:
            values_and_rows.append((row, value))
    if not values_and_rows:
        return None
    values = [value for _row, value in values_and_rows]
    label = value_col if op != "count" else "rows"
    text = f"{_op_label(op)} {label}: {_format_number(_apply_aggregate(op, values))}"
    return text, _refs_for_rows([row for row, _value in values_and_rows], value_col or "")


def _rank(
    query: str, rows: list[_StructuredRow], max_rows: int
) -> tuple[str, list[_CellRef]] | None:
    value_col = _best_numeric_column(query, rows)
    if not value_col:
        return None
    ascending = bool(re.search(r"\b(lowest|least|bottom)\b|下位|最小|最も少ない", query, re.I))
    top_n_match = _TOP_N.search(query)
    top_n = min(max_rows, int(top_n_match.group(1)) if top_n_match else max_rows)
    ranked = [
        (row, _numeric(row.values.get(value_col, "")))
        for row in rows
        if _numeric(row.values.get(value_col, "")) is not None
    ]
    ranked.sort(key=lambda item: item[1] or 0.0, reverse=not ascending)
    selected = ranked[:top_n]
    if not selected:
        return None
    label_col = _label_column(rows, exclude=value_col)
    parts = [
        f"{row.values.get(label_col, row.row_id)}={_format_number(value or 0.0)}"
        for row, value in selected
    ]
    return f"{'Lowest' if ascending else 'Highest'} {value_col}: " + "; ".join(
        parts
    ), _refs_for_rows([row for row, _value in selected], value_col)


def _latest(query: str, rows: list[_StructuredRow]) -> tuple[str, list[_CellRef]] | None:
    date_col = _date_column(rows)
    if not date_col:
        return None
    dated = [(row, _date_key(row.values.get(date_col, ""))) for row in rows]
    dated = [(row, key) for row, key in dated if key]
    if not dated:
        return None
    row = max(dated, key=lambda item: item[1])[0]
    target_col = _best_column(query, rows, exclude=date_col) or _label_column(
        rows, exclude=date_col
    )
    if not target_col:
        return None
    text = f"Latest {target_col}: {row.values.get(target_col, '')} ({date_col}: {row.values.get(date_col, '')})"
    return text, _refs_for_rows([row], target_col)


def _numeric_filter(
    query: str, rows: list[_StructuredRow], max_rows: int
) -> tuple[str, list[_CellRef]] | None:
    value_col = _best_numeric_column(query, rows)
    threshold = _query_number(query)
    if not value_col or threshold is None:
        return None
    if re.search(r"\b(less|below|under)\b|<|以下|未満|より小さい", query, re.I):
        predicate = lambda value: value < threshold
        label = f"{value_col} < {_format_number(threshold)}"
    else:
        predicate = lambda value: value > threshold
        label = f"{value_col} > {_format_number(threshold)}"
    matches = [
        row
        for row in rows
        if (value := _numeric(row.values.get(value_col, ""))) is not None and predicate(value)
    ][:max_rows]
    if not matches:
        return None
    label_col = _label_column(rows, exclude=value_col)
    text = (
        label
        + ": "
        + "; ".join(
            f"{row.values.get(label_col, row.row_id)}={row.values.get(value_col, '')}"
            for row in matches
        )
    )
    return text, _refs_for_rows(matches, value_col)


def _period_filter(
    query: str, rows: list[_StructuredRow], max_rows: int
) -> tuple[str, list[_CellRef]] | None:
    date_col = _date_column(rows)
    if not date_col:
        return None
    lower = _query_date_lower_bound(query)
    if not lower:
        return None
    matches = [
        row for row in rows if (key := _date_key(row.values.get(date_col, ""))) and key >= lower
    ][:max_rows]
    if not matches:
        return None
    label_col = _label_column(rows, exclude=date_col)
    text = f"Rows since {lower}: " + "; ".join(
        f"{row.values.get(label_col, row.row_id)} ({date_col}: {row.values.get(date_col, '')})"
        for row in matches
    )
    return text, _refs_for_rows(matches, date_col)


def _aggregate_op(query: str) -> str:
    q = query.lower()
    if "average" in q or re.search(r"\bavg\b", q) or "平均" in query:
        return "average"
    if "median" in q or "中央値" in query:
        return "median"
    if "total" in q or "sum" in q or "合計" in query:
        return "sum"
    if "count" in q or "件数" in query or "何件" in query:
        return "count"
    return "sum"


def _op_label(op: str) -> str:
    return {"sum": "Total", "average": "Average", "median": "Median", "count": "Count"}[op]


def _apply_aggregate(op: str, values: list[float]) -> float:
    if op == "average":
        return sum(values) / len(values)
    if op == "median":
        return float(median(values))
    if op == "count":
        return float(len(values))
    return sum(values)


def _best_numeric_column(query: str, rows: list[_StructuredRow]) -> str:
    candidates = [
        column
        for column in _columns(rows)
        if any(_numeric(row.values.get(column, "")) is not None for row in rows)
    ]
    return _best_named_column(query, candidates)


def _best_column(query: str, rows: list[_StructuredRow], *, exclude: str = "") -> str:
    candidates = [column for column in _columns(rows) if column != exclude]
    return _best_named_column(query, candidates)


def _best_named_column(query: str, candidates: list[str], *, fallback: bool = True) -> str:
    if not candidates:
        return ""
    terms = _terms(query)
    scored = [(len(terms & _terms(column)), len(_terms(column)), column) for column in candidates]
    scored.sort(key=lambda item: (item[0], item[1]), reverse=True)
    return scored[0][2] if scored[0][0] > 0 else candidates[0] if fallback else ""


def _group_column(query: str, rows: list[_StructuredRow], *, exclude: str = "") -> str:
    by_match = re.search(r"\b(?:by|per)\s+([A-Za-z0-9_一-龯ぁ-んァ-ン ]+)", query, re.I)
    if by_match:
        wanted = by_match.group(1)
        column = _best_named_column(
            wanted, [c for c in _columns(rows) if c != exclude], fallback=False
        )
        if column:
            return column
    non_numeric = [
        column
        for column in _columns(rows)
        if column != exclude
        and not any(_numeric(row.values.get(column, "")) is not None for row in rows)
    ]
    return _best_named_column(query, non_numeric, fallback=False)


def _label_column(rows: list[_StructuredRow], *, exclude: str = "") -> str:
    for column in _columns(rows):
        if column != exclude and not any(
            _numeric(row.values.get(column, "")) is not None for row in rows
        ):
            return column
    return next((column for column in _columns(rows) if column != exclude), "")


def _date_column(rows: list[_StructuredRow]) -> str:
    candidates = [
        column
        for column in _columns(rows)
        if any(_date_key(row.values.get(column, "")) for row in rows)
    ]
    if not candidates:
        return ""
    preferred = [
        c
        for c in candidates
        if any(t in c.lower() for t in ("date", "month", "updated", "created", "settled"))
    ]
    return (preferred or candidates)[0]


def _columns(rows: list[_StructuredRow]) -> list[str]:
    seen: list[str] = []
    for row in rows:
        for column in row.values:
            if column not in seen:
                seen.append(column)
    return seen


def _refs_for_rows(rows: list[_StructuredRow], column: str) -> list[_CellRef]:
    refs: list[_CellRef] = []
    for row in rows:
        cell = row.cells.get(column) or next(iter(row.cells.values()), "")
        refs.append(
            _CellRef(
                document_id=row.document_id,
                source_id=row.source_id,
                version=row.version,
                indexed_at=row.indexed_at,
                sheet_name=row.sheet_name,
                row_id=row.row_id,
                cell_range=cell,
            )
        )
    return refs


def _numeric(value: str | None) -> float | None:
    if not value:
        return None
    match = _NUMBER.search(str(value).replace("%", ""))
    if not match:
        return None
    try:
        return float(match.group(0).replace(",", ""))
    except ValueError:
        return None


def _query_number(query: str) -> float | None:
    match = _NUMBER.search(query)
    if not match:
        return None
    return float(match.group(0).replace(",", ""))


def _date_key(value: str | None) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    ym = _YEAR_MONTH.search(raw)
    if ym:
        return f"{int(ym.group(1)):04d}-{int(ym.group(2)):02d}-00"
    for token, month in _MONTHS.items():
        if re.search(rf"\b{token}\b", raw, re.I):
            year_match = _YEAR.search(raw)
            year = int(year_match.group(1)) if year_match else 0
            return f"{year:04d}-{month:02d}-00"
    try:
        return datetime.fromisoformat(raw.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        return ""


def _query_date_lower_bound(query: str) -> str:
    ym = _YEAR_MONTH.search(query)
    if ym:
        return f"{int(ym.group(1)):04d}-{int(ym.group(2)):02d}-00"
    year = _YEAR.search(query)
    if year:
        return f"{int(year.group(1)):04d}-00-00"
    return ""


def _format_number(value: float) -> str:
    if value.is_integer():
        return str(int(value))
    return f"{value:.3f}".rstrip("0").rstrip(".")


def _terms(text: str) -> set[str]:
    return {token.lower() for token in _WORD.findall(text or "") if len(token) > 1}
