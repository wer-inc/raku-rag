"""T040 — Parser (PDF/Markdown/HTML/text). MVP implements text/markdown/html + NFKC normalize.

日本語対応 (FR-003a): NFKC 正規化。PDF は production で pdf parser を差し替え（同 interface）。

US2 (T025/T026) ADDS, behind the SAME 001 ``Parser`` abstraction (no new search/chunk mechanism):
  - ``DocxParser``     — DOCX paragraphs + table cells (python-docx) -> normalized text.
  - ``SpreadsheetParser`` — XLSX (openpyxl) + CSV (stdlib) -> normalized text where EACH data cell is
    its own ``\\n\\n``-delimited block carrying a ``sheet!R{row}C{col}`` cell anchor (header row = R1,
    1-based row/col). The SentenceChunker then yields one chunk per cell, so a citation resolves to a
    single cell (FR-MFG-002). The anchor token is parsed back out of the chunk text / Chunk.metadata
    by the manufacturing search overlay to expose the cell coordinate.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from typing import Sequence

from raku_rag.interfaces.base import Parser

_TAG = re.compile(r"<[^>]+>")
_SUPPORTED = {"text/plain", "text/markdown", "text/html"}

# OOXML content types (US2). Mirrored in tests/manufacturing/test_*formats / test_cell_citation.
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_CONTENT_TYPES = {"text/csv", "application/csv"}

# sheet!R{row}C{col} cell anchor. header is row 1; rows/cols are 1-based.
_CELL_ANCHOR = re.compile(r"([A-Za-z0-9_]+)!R(\d+)C(\d+)")


def cell_anchor(sheet: str, row: int, col: int) -> str:
    """Build the canonical ``sheet!R{row}C{col}`` cell anchor token (FR-MFG-002)."""
    return f"{sheet}!R{row}C{col}"


def parse_cell_anchor(text: str) -> tuple[str, int, int] | None:
    """Extract ``(sheet, row, col)`` from the first ``sheet!R{row}C{col}`` token in ``text``."""
    m = _CELL_ANCHOR.search(text or "")
    if not m:
        return None
    return m.group(1), int(m.group(2)), int(m.group(3))


def read_xlsx_sheets(raw: bytes) -> list[tuple[str, list[list[str]]]]:
    """(sheet_name, rows-as-strings) for every worksheet. Shared by SpreadsheetParser + structured."""
    from openpyxl import load_workbook  # local import: optional [manufacturing] dep

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheets: list[tuple[str, list[list[str]]]] = []
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else str(v) for v in row])
        sheets.append((ws.title, rows))
    wb.close()
    return sheets


def read_csv_sheets(raw: bytes) -> list[tuple[str, list[list[str]]]]:
    """CSV rows under a single stable logical sheet name (CSV has no sheet)."""
    text = raw.decode("utf-8", errors="replace")
    rows = [list(r) for r in csv.reader(io.StringIO(text))]
    return [("sheet1", rows)]


def read_spreadsheet_sheets(raw: bytes, content_type: str) -> list[tuple[str, list[list[str]]]]:
    """Dispatch XLSX/CSV to the matching reader."""
    if content_type == XLSX_CONTENT_TYPE:
        return read_xlsx_sheets(raw)
    return read_csv_sheets(raw)


class TextParser(Parser):
    def supports(self, content_type: str) -> bool:
        return content_type in _SUPPORTED

    def parse(self, raw: bytes, content_type: str) -> str:
        text = raw.decode("utf-8", errors="replace")
        if content_type == "text/html":
            text = _TAG.sub(" ", text)
        text = unicodedata.normalize("NFKC", text)
        # collapse excess whitespace but keep line/paragraph boundaries
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()


class DocxParser(Parser):
    """T025 — DOCX -> normalized text (python-docx), behind the 001 ``Parser`` abstraction.

    Extracts BOTH body paragraphs AND table-cell content so a citation can be grounded in table
    data (FR-MFG-001). Each paragraph and each table row is emitted as its own ``\\n\\n`` block so the
    001 SentenceChunker keeps them as distinct, citable chunks. NFKC-normalized like TextParser.
    """

    def supports(self, content_type: str) -> bool:
        return content_type == DOCX_CONTENT_TYPE

    def parse(self, raw: bytes, content_type: str) -> str:
        from docx import Document as _DocxDocument  # local import: optional [manufacturing] dep

        doc = _DocxDocument(io.BytesIO(raw))
        blocks: list[str] = []
        for para in doc.paragraphs:
            text = (para.text or "").strip()
            if text:
                blocks.append(text)
        for table in doc.tables:
            for row in table.rows:
                cells = [(c.text or "").strip() for c in row.cells]
                line = " | ".join(cell for cell in cells if cell)
                if line:
                    blocks.append(line)
        normalized = [unicodedata.normalize("NFKC", b) for b in blocks]
        normalized = [re.sub(r"[ \t]+", " ", b).strip() for b in normalized]
        return "\n\n".join(b for b in normalized if b)


class SpreadsheetParser(Parser):
    """T026 — XLSX (openpyxl) + CSV (stdlib) -> normalized text with per-cell ``sheet!R{row}C{col}``
    anchors (FR-MFG-002), behind the 001 ``Parser`` abstraction.

    Layout (header row = R1, rows/cols 1-based): every NON-EMPTY data cell (rows below the header)
    becomes its own ``\\n\\n`` block of the form ``"{sheet}!R{row}C{col} {header}: {value}"``. The
    header text is included so a natural-language query (e.g. "remedy …") still retrieves the right
    cell, while the leading anchor token preserves the exact coordinate through the 001
    chunk/citation/offset path. One cell per block => the SentenceChunker yields one chunk per cell,
    so distinct cells resolve to distinct anchors.
    """

    def supports(self, content_type: str) -> bool:
        return content_type == XLSX_CONTENT_TYPE or content_type in CSV_CONTENT_TYPES

    def parse(self, raw: bytes, content_type: str) -> str:
        sheets = read_spreadsheet_sheets(raw, content_type)
        return "\n\n".join(self._emit_cells(sheets))

    # --- emit ------------------------------------------------------------------------------------
    def _emit_cells(self, sheets: list[tuple[str, list[list[str]]]]):
        for sheet, rows in sheets:
            if not rows:
                continue
            header = rows[0]
            # Data rows are rows below the header (R1 = header, first data row = R2).
            for r_idx, row in enumerate(rows[1:], start=2):
                for c_idx, value in enumerate(row, start=1):
                    cell = unicodedata.normalize("NFKC", str(value)).strip()
                    cell = re.sub(r"[ \t]+", " ", cell)
                    if not cell:
                        continue
                    col_name = ""
                    if c_idx - 1 < len(header):
                        col_name = unicodedata.normalize("NFKC", str(header[c_idx - 1])).strip()
                    anchor = cell_anchor(sheet, r_idx, c_idx)
                    label = f"{col_name}: " if col_name else ""
                    yield f"{anchor} {label}{cell}"


class CompositeParser(Parser):
    """Dispatch a parse to the first delegate that ``supports`` the content type.

    Lets the manufacturing ingestion path reuse the 001 ``IngestionService`` (which holds a single
    ``Parser``) while accepting text/markdown/html + DOCX + XLSX/CSV. ``supports`` is the union of the
    delegates' support, so ``IngestionService`` still rejects genuinely unsupported types (FR-030).
    """

    def __init__(self, parsers: Sequence[Parser]) -> None:
        self._parsers = tuple(parsers)

    def supports(self, content_type: str) -> bool:
        return any(p.supports(content_type) for p in self._parsers)

    def parse(self, raw: bytes, content_type: str) -> str:
        for p in self._parsers:
            if p.supports(content_type):
                return p.parse(raw, content_type)
        raise ValueError(f"unsupported content_type: {content_type}")
