"""T067 — Quickstart S1-S11 validation (quickstart.md; spec FR-MFG-027/028).

Maps EACH quickstart validation scenario (S1 through S11) to an executable assertion against the
in-memory ``ManufacturingSystem`` composition root. This is a VALIDATION harness, not a redefinition
of behaviour: every scenario exercises an already-built entrypoint and asserts the quickstart's stated
acceptance criterion. The deep correctness of each safety hard gate is mechanism-pinned by its own
dedicated test (referenced per scenario); here we pin that the quickstart-level acceptance criteria
hold when reached through the public entrypoints, end to end.

Scenario -> entrypoint(s) -> acceptance criterion:
  S1  ingest_manufacturing_file + search        — mfg ingest + metadata filter + approval tags (FR-MFG-001/002/003/004)
  S2  answer                                     — grounded immediate answer w/ approved citation (FR-MFG-005)
  S3  answer (high-risk)                         — approved-citation-required safety gate (SC-MFG-006)
  S4  answer (obsolete/draft evidence)           — obsolete/draft never primary (SC-MFG-011)
  S5  search_trouble_cases                       — similar cases, provisional/permanent, candidate (Hard Rule 4)
  S6  generate_draft/assign_reviewer/review_draft— draft-only, reviewer-approved (SC-MFG-007)
  S7  grant_scope + search/answer                — department/factory/role ACL mapping (SC-MFG-008)
  S8  get_data_use_policy/update.../capability   — no-train default + opt-in invariant + block (SC-MFG-009)
  S9  export_audit + verify_chain                — audit coverage + reference-only + tamper-evidence (SC-MFG-010)
  S10 safety_telemetry/dashboard/kpi             — telemetry breakdown + KPI export (SC-MFG-012/013)
  S11 the full vertical slice                    — S1-S10 invariants compose end-to-end (FR-MFG-027)

The PDF / visual-RAG bullet of S1/S11 is part of the 001 base platform (001 quickstart) — there is no
manufacturing-layer PDF parser to exercise here, so this harness validates the DOCX/XLSX/CSV+text
ingestion the 002 layer actually owns and notes the PDF path as 001-base scope.

stdlib only (+ python-docx/openpyxl for the DOCX/XLSX fixtures, already present). Authoritative:
quickstart.md S1-S11; spec FR-MFG-027/028, SC-MFG-006~013. Assertion style mirrors the dedicated
per-scenario tests under tests/manufacturing/.
"""

from __future__ import annotations

import csv
import io
import tempfile
import unittest
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import Workbook

from raku_rag.domain.models import Citation, ScopeType, SubjectType
from raku_rag.manufacturing.domain.acl_mapping import ManufacturingScope, factory_scope
from raku_rag.manufacturing.domain.draft import CreatedBy, DraftStatus, DraftType
from raku_rag.manufacturing.domain.entities import (
    Countermeasure,
    CountermeasureType,
    FailureMode,
    MeasureClass,
    TroubleCase,
)
from raku_rag.manufacturing.domain.metadata import (
    ApprovalSource,
    ApprovalStatus,
    DocumentKind,
    ManufacturingDocumentMetadata,
)
from raku_rag.manufacturing.domain.policy import NoTrainFallback
from raku_rag.manufacturing.domain.safety import SafetyBlockReason
from tests.manufacturing.helpers import T, claims, fresh, mfg_meta

DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_CT = "text/csv"


def _val(x):
    return getattr(x, "value", x)


def _attr_or_key(obj, name):
    if isinstance(obj, dict):
        return obj.get(name)
    return getattr(obj, name, None)


# ------------------------------------------------------------------------------------------------
# S1. 製造業文書の取り込み + メタデータ + 承認 (US2; FR-MFG-001/002/003/004).
# ------------------------------------------------------------------------------------------------
class S1_IngestMetadataApproval(unittest.TestCase):
    def setUp(self) -> None:
        self.sys = fresh()
        self.sys.grant(T, ScopeType.COLLECTION, "c", SubjectType.USER, "op")
        self.op = claims(T, "op")
        self._tmp = tempfile.TemporaryDirectory()
        tmp = Path(self._tmp.name)
        self.xlsx = tmp / "ledger.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "ledger"
        ws.append(["equipment", "alarm", "remedy"])
        ws.append(["pump17", "E152", "replace the impeller seal on pump17 for alarm E152"])
        wb.save(str(self.xlsx))

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_s1_ingest_succeeds_and_search_returns_approval_and_cell_anchor(self) -> None:
        job = self.sys.ingest_manufacturing_file(
            tenant_id=T,
            collection_id="c",
            document_id="ledger1",
            path=str(self.xlsx),
            content_type=XLSX_CT,
            metadata=ManufacturingDocumentMetadata(
                tenant_id=T,
                document_id="ledger1",
                approval_status=ApprovalStatus.APPROVED,
                effective_date="2026-01-10",
                document_kind=DocumentKind.LEDGER,
                equipment="pump17",
                alarm_code="E152",
            ),
        )
        # S1.1/S1.2: ingestion succeeds via the reused 001 path (job succeeded, chunks indexed).
        self.assertEqual(job.status, "succeeded", getattr(job, "failure_reason", ""))
        self.assertGreater(job.chunk_count, 0)
        # S1.3: search filters by manufacturing metadata and returns approval_status/effective_date.
        results = self.sys.search(
            self.op, "impeller seal pump17 alarm E152", manufacturing_filters={"alarm_code": "E152"}
        )
        hit = next((r for r in results if r.document_id == "ledger1"), None)
        self.assertIsNotNone(hit, "metadata-filtered search must return the ingested ledger")
        self.assertEqual(hit.approval_status, "approved")
        self.assertEqual(hit.effective_date, "2026-01-10")
        # S1.3: an XLSX-derived citation carries the sheet!RxCy cell coordinate (FR-MFG-002).
        cell_hit = next((r for r in results if r.cell_anchor is not None), None)
        self.assertIsNotNone(
            cell_hit, "a spreadsheet-derived result must carry a cell coordinate (FR-MFG-002)"
        )
        self.assertIsNotNone(cell_hit.sheet)
        self.assertIsNotNone(cell_hit.row)
        self.assertIsNotNone(cell_hit.col)


# ------------------------------------------------------------------------------------------------
# S2. 根拠付き即答 (US1; FR-MFG-005).
# ------------------------------------------------------------------------------------------------
class S2_GroundedImmediateAnswer(unittest.TestCase):
    def setUp(self) -> None:
        self.sys = fresh()
        self.sys.ingest_manufacturing(
            tenant_id=T,
            collection_id="c",
            document_id="manual1",
            text="The conveyor motor lubrication interval is every ninety days under normal load.",
            metadata=mfg_meta(tenant_id=T, document_id="manual1"),
        )
        self.sys.grant(T, ScopeType.COLLECTION, "c", SubjectType.USER, "op")
        self.op = claims(T, "op")

    def test_s2_approved_manual_answers_with_citation_and_used_chunks(self) -> None:
        ans = self.sys.answer(self.op, "what is the conveyor motor lubrication interval?")
        self.assertEqual(ans.status, "ok")
        self.assertTrue(ans.text)
        self.assertTrue(ans.citations, "S2: citation with document/chunk required")
        self.assertTrue(ans.used_chunks, "S2: used_chunks required")
        self.assertEqual(ans.citations[0].approval_status, "approved")
        self.assertEqual(ans.citations[0].effective_date, "2026-01-10")


# ------------------------------------------------------------------------------------------------
# S3. High-risk safety gate — approved 引用必須 (SC-MFG-006; FR-MFG-005/007/015).
# ------------------------------------------------------------------------------------------------
class S3_HighRiskApprovedCitationRequired(unittest.TestCase):
    def _seed(self, status: ApprovalStatus, effective: str | None) -> None:
        self.sys.ingest_manufacturing(
            tenant_id=T,
            collection_id="c",
            document_id="lockout",
            text="To disassemble the press, stop the machine, apply lockout tagout, release pressure.",
            metadata=mfg_meta(
                tenant_id=T,
                document_id="lockout",
                approval_status=status,
                effective_date=effective,
                document_kind=DocumentKind.WORK_INSTRUCTION,
                safety_category="lockout_tagout",
                hazard_tags=("設備停止", "分解", "高圧"),
            ),
        )
        self.sys.grant(T, ScopeType.COLLECTION, "c", SubjectType.USER, "op")

    def setUp(self) -> None:
        self.sys = fresh()
        self.op = claims(T, "op")

    def test_s3_1_no_approved_evidence_blocks_with_reason_and_no_assertion(self) -> None:
        self._seed(ApprovalStatus.DRAFT, None)  # only draft evidence => not approved+effective
        ans = self.sys.answer(self.op, "How do I disassemble the press safely?")
        self.assertTrue(ans.high_risk)
        self.assertEqual(ans.status, "insufficient_evidence")
        self.assertEqual(ans.safety_block_reason, SafetyBlockReason.APPROVED_CITATION_MISSING.value)
        self.assertFalse(ans.text, "S3.1: must not return a guessed procedure")

    def test_s3_2_ambiguous_safety_query_is_fail_safe_high_risk(self) -> None:
        self.sys.ingest_manufacturing(
            tenant_id=T,
            collection_id="c",
            document_id="bland",
            text="The procedure document covers the handling step for the unit on the line.",
            metadata=mfg_meta(tenant_id=T, document_id="bland"),
        )
        self.sys.grant(T, ScopeType.COLLECTION, "c", SubjectType.USER, "op")
        ans = self.sys.answer(self.op, "how should I handle this?", intent_hint="ambiguous")
        self.assertTrue(
            ans.high_risk, "S3.2: ambiguous safety query is treated high-risk (fail-safe)"
        )

    def test_s3_3_high_risk_with_approved_citation_answers_and_requires_onsite(self) -> None:
        self._seed(ApprovalStatus.APPROVED, "2026-01-10")
        ans = self.sys.answer(self.op, "How do I disassemble the press safely?")
        self.assertTrue(ans.high_risk)
        self.assertEqual(ans.status, "ok")
        self.assertEqual(ans.citations[0].approval_status, "approved")
        self.assertTrue(
            ans.requires_onsite_confirmation, "S3.3: on-site confirmation required (FR-MFG-007)"
        )


# ------------------------------------------------------------------------------------------------
# S4. obsolete / draft の扱い (SC-MFG-011; FR-MFG-006).
# ------------------------------------------------------------------------------------------------
class S4_ObsoleteDraftHandling(unittest.TestCase):
    def setUp(self) -> None:
        self.sys = fresh()
        self.op = claims(T, "op")

    def test_s4_1_obsolete_only_evidence_not_primary_and_warns(self) -> None:
        self.sys.ingest_manufacturing(
            tenant_id=T,
            collection_id="c",
            document_id="old_spec",
            text="The torque specification for the flange bolt is forty newton meters.",
            metadata=mfg_meta(
                tenant_id=T,
                document_id="old_spec",
                approval_status=ApprovalStatus.OBSOLETE,
                effective_date="2020-01-01",
                obsolete_at="2025-12-31",
            ),
        )
        self.sys.grant(T, ScopeType.COLLECTION, "c", SubjectType.USER, "op")
        ans = self.sys.answer(self.op, "what is the torque specification for the flange bolt?")
        self.assertNotEqual(ans.status, "ok", "S4.1: obsolete-only evidence must not assert")
        self.assertTrue(ans.obsolete_warning, "S4.1: obsolete reference requires obsolete_warning")

    def test_s4_2_draft_only_evidence_not_formal_basis(self) -> None:
        self.sys.ingest_manufacturing(
            tenant_id=T,
            collection_id="c",
            document_id="draft_spec",
            text="The torque specification for the flange bolt is forty newton meters.",
            metadata=mfg_meta(
                tenant_id=T,
                document_id="draft_spec",
                approval_status=ApprovalStatus.DRAFT,
                effective_date=None,
            ),
        )
        self.sys.grant(T, ScopeType.COLLECTION, "c", SubjectType.USER, "op")
        ans = self.sys.answer(self.op, "what is the torque specification for the flange bolt?")
        self.assertNotEqual(
            ans.status, "ok", "S4.2: draft-only evidence must not be a formal basis"
        )
        for c in ans.citations:
            self.assertNotEqual(getattr(c, "approval_status", None), "approved")


# ------------------------------------------------------------------------------------------------
# S5. 類似トラブル事例 (US3; FR-MFG-008/009, Hard Rule 4).
# ------------------------------------------------------------------------------------------------
class S5_SimilarTroubleCases(unittest.TestCase):
    def setUp(self) -> None:
        self.sys = fresh()
        _, coll = factory_scope("facA")
        self.coll = coll
        self.sys.register_trouble_case(
            tenant_id=T,
            collection_id=coll,
            source_document_id="tr_A",
            text=(
                "Trouble report: the gearbox showed vibration increase and abnormal noise. Root cause: "
                "bearing wear. Provisional countermeasure: reduce feed rate. Permanent countermeasure: "
                "replace the worn bearing. Recurrence prevention: add the bearing to the schedule."
            ),
            metadata=mfg_meta(tenant_id=T, document_id="tr_A"),
            trouble_case=TroubleCase(
                tenant_id=T,
                trouble_case_id="tc_A",
                symptom="振動増加＋異音",
                equipment_id="eq_gearbox",
                process_id="pr_assembly",
                failure_mode_id="fm_bearing",
                source_document_id="tr_A",
            ),
            failure_mode=FailureMode(
                tenant_id=T,
                failure_mode_id="fm_bearing",
                name="軸受摩耗",
                description="bearing wear",
            ),
            countermeasures=(
                Countermeasure(
                    tenant_id=T,
                    measure_id="m_prov",
                    trouble_case_id="tc_A",
                    description="reduce feed rate",
                    measure_class=MeasureClass.PROVISIONAL,
                    source_document_id="tr_A",
                ),
                Countermeasure(
                    tenant_id=T,
                    measure_id="m_perm",
                    trouble_case_id="tc_A",
                    description="replace the worn bearing",
                    measure_class=MeasureClass.PERMANENT,
                    source_document_id="tr_A",
                ),
            ),
            recurrence_prevention="add the bearing to the periodic replacement schedule",
        )
        self.sys.grant_scope(
            ManufacturingScope(
                tenant_id=T,
                department="maintenance_dept",
                roles=("technician",),
                factory_ids=("facA",),
            )
        )
        self.op = claims(T, "fa_user", groups=["maintenance_dept"], roles=["technician"])

    def test_s5_similar_case_with_cause_split_measures_and_candidate_label(self) -> None:
        resp = self.sys.search_trouble_cases(
            self.op, "振動増加 異音 vibration increase abnormal noise gearbox bearing"
        )
        self.assertEqual(_val(resp.status), "ok")
        m = next((x for x in resp.results if x.trouble_case_id == "tc_A"), None)
        self.assertIsNotNone(m, "S5: the similar gearbox case must be listed")
        self.assertIsNotNone(m.failure_mode, "S5: cause (FailureMode) accompanies the case")
        self.assertTrue(m.recurrence_prevention, "S5: recurrence prevention accompanies the case")
        self.assertTrue(m.citations, "S5: the case is cited (出典つき)")
        # S5: provisional/permanent split.
        self.assertIn("m_prov", {c.measure_id for c in m.countermeasures.provisional})
        self.assertIn("m_perm", {c.measure_id for c in m.countermeasures.permanent})
        # S5 / Hard Rule 4: even the PERMANENT past-case measure displays as candidate/past-example.
        perm = next(c for c in m.countermeasures.permanent if c.measure_id == "m_perm")
        self.assertEqual(_val(perm.type), CountermeasureType.CANDIDATE.value)
        self.assertTrue(perm.label, "S5: a past-example candidate label is shown")


# ------------------------------------------------------------------------------------------------
# S6. ドラフト生成 (US4; SC-MFG-007, G1).
# ------------------------------------------------------------------------------------------------
class S6_DraftGeneration(unittest.TestCase):
    def setUp(self) -> None:
        self.sys = fresh()
        self.author = claims(T, "author")

    def test_s6_1_checklist_draft_has_status_draft_and_source_citations(self) -> None:
        art = self.sys.generate_draft(
            principal=self.author,
            kind=DraftType.CHECKLIST,
            context_citations=(
                Citation(
                    kind="text",
                    document_id="d1",
                    source_id="src",
                    version=1,
                    retrieval_score=0.9,
                    chunk_id="d1#0",
                ),
            ),
            source_document_ids=("d1",),
        )
        self.assertEqual(
            _val(art.status), DraftStatus.DRAFT.value, "S6.1: kind=checklist is status=draft"
        )
        self.assertTrue(art.source_citations, "S6.1: source_citations present")
        self.assertEqual(_val(art.created_by), CreatedBy.AI.value)

    def test_s6_3_faq_not_auto_approved(self) -> None:
        faq = self.sys.generate_draft(principal=self.author, kind=DraftType.FAQ)
        self.assertEqual(
            _val(faq.status), DraftStatus.DRAFT.value, "S6.3: FAQ is draft, never auto-approved"
        )

    def test_s6_4_reviewer_only_confirms_via_assign_then_review(self) -> None:
        art = self.sys.generate_draft(principal=self.author, kind=DraftType.CHECKLIST)
        self.sys.assign_reviewer(tenant_id=T, artifact_id=art.artifact_id, reviewer_id="rev_1")
        approved = self.sys.review_draft(
            tenant_id=T,
            artifact_id=art.artifact_id,
            reviewer=claims(T, "rev_1", roles=("reviewer",)),
            decision="approved",
        )
        self.assertEqual(
            _val(approved.status), DraftStatus.APPROVED.value, "S6.4: reviewer confirms"
        )
        self.assertEqual(approved.reviewer_id, "rev_1")

    def test_s6_5_ai_self_approve_rejected(self) -> None:
        art = self.sys.generate_draft(principal=self.author, kind=DraftType.CHECKLIST)
        leaked = False
        try:
            r = self.sys.review_draft(
                tenant_id=T, artifact_id=art.artifact_id, reviewer=None, decision="approved"
            )
            leaked = _val(r.status) == DraftStatus.APPROVED.value
        except (ValueError, PermissionError, TypeError):
            pass
        self.assertFalse(leaked, "S6.5: created_by=ai cannot self-approve (SC-MFG-007 = 0)")
        self.assertEqual(
            _val(self.sys.get_draft(T, art.artifact_id).status), DraftStatus.DRAFT.value
        )


# ------------------------------------------------------------------------------------------------
# S7. ACL マッピング (US6; SC-MFG-008, 001 ACL reuse).
# ------------------------------------------------------------------------------------------------
class S7_AclMapping(unittest.TestCase):
    def setUp(self) -> None:
        self.sys = fresh()
        _, self.coll_b = factory_scope("facB")
        # A confidential factory-B / quality document.
        self.sys.ingest_manufacturing(
            tenant_id=T,
            collection_id=self.coll_b,
            document_id="conf_B",
            text="Confidential: customer AcmeMotors cathode line cathode-delamination defect on drawing DRW-9.",
            metadata=mfg_meta(
                tenant_id=T,
                document_id="conf_B",
                customer="AcmeMotors",
                defect_type="cathode-delamination",
            ),
        )
        # Authorized factory-B quality user; unauthorized factory-A maintenance user.
        self.sys.grant_scope(
            ManufacturingScope(
                tenant_id=T, department="quality_dept", roles=("supervisor",), factory_ids=("facB",)
            )
        )
        self.sys.grant_scope(
            ManufacturingScope(
                tenant_id=T,
                department="maintenance_dept",
                roles=("technician",),
                factory_ids=("facA",),
            )
        )
        self.authorized = claims(T, "qb", groups=["quality_dept"], roles=["supervisor"])
        self.unauthorized = claims(T, "fa", groups=["maintenance_dept"], roles=["technician"])
        self.probe = "AcmeMotors cathode-delamination defect drawing DRW-9 customer"

    def test_s7_unauthorized_user_sees_no_confidential_doc_or_citation(self) -> None:
        results = self.sys.search(self.unauthorized, self.probe)
        self.assertTrue(
            all(r.document_id != "conf_B" for r in results),
            "S7: confidential doc must not leak into search",
        )
        ans = self.sys.answer(
            self.unauthorized,
            "what is the AcmeMotors cathode-delamination defect?",
            collection_id=self.coll_b,
        )
        for c in ans.citations:
            self.assertNotEqual(
                c.document_id,
                "conf_B",
                "S7: confidential doc must not be cited to unauthorized user",
            )

    def test_s7_authorized_user_can_see_it_positive_control(self) -> None:
        results = self.sys.search(self.authorized, self.probe)
        self.assertTrue(
            any(r.document_id == "conf_B" for r in results),
            "S7: authorized user must find their own doc",
        )


# ------------------------------------------------------------------------------------------------
# S8. No-train / Governance (FR-MFG-016~020/029; SC-MFG-009, GQ1).
# ------------------------------------------------------------------------------------------------
class S8_NoTrainGovernance(unittest.TestCase):
    def setUp(self) -> None:
        self.admin = claims(T, "admin", roles=("admin",))

    def test_s8_1_default_data_use_policy(self) -> None:
        sys = fresh()
        policy = sys.get_data_use_policy(T)
        self.assertTrue(policy.no_train_default)
        self.assertFalse(policy.training_opt_in)
        self.assertEqual(policy.no_train_fallback, NoTrainFallback.BLOCK)
        self.assertEqual(policy.retention_customer, 365)
        self.assertEqual(policy.retention_audit, 365)

    def test_s8_2_training_without_optin_is_refused(self) -> None:
        sys = fresh()
        with self.assertRaises(
            Exception, msg="S8.2: training without opt-in must be refused (SC-MFG-009 = 0)"
        ):
            sys.use_for_training(tenant_id=T, data_kind="answer", actor=self.admin)

    def test_s8_3_capability_without_no_train_provider_is_blocked(self) -> None:
        # A capability whose only provider is NOT no-train-guaranteed => temporarily_unavailable (GQ1).
        from raku_rag.manufacturing.app import ManufacturingSystem

        sys = ManufacturingSystem(
            provider_capabilities={"answer_llm": ("untrusted_vendor",)},
            no_train_providers=(),
        )
        self.assertEqual(
            sys.capability_status(T, "answer_llm"),
            "temporarily_unavailable",
            "S8.3: a non-no-train capability must be blocked, not silently degraded (GQ1)",
        )

    def test_s8_4_optin_without_contract_ref_rejected_and_version_bumps_on_valid_change(
        self,
    ) -> None:
        sys = fresh()
        with self.assertRaises(
            Exception, msg="S8.4: training_opt_in=True without opt_in_contract_ref must be rejected"
        ):
            sys.update_data_use_policy(
                tenant_id=T, patch={"training_opt_in": True}, actor=self.admin
            )
        before = sys.get_data_use_policy(T).policy_version
        sys.update_data_use_policy(tenant_id=T, patch={"retention_customer": 730}, actor=self.admin)
        after = sys.get_data_use_policy(T).policy_version
        self.assertNotEqual(before, after, "S8.4: a policy change must bump policy_version")


# ------------------------------------------------------------------------------------------------
# S9. Audit coverage (FR-MFG-021~023; SC-MFG-010).
# ------------------------------------------------------------------------------------------------
class S9_AuditCoverage(unittest.TestCase):
    def setUp(self) -> None:
        self.sys = fresh()
        self.admin = claims(T, "admin", roles=("admin",))
        self.op = claims(T, "op")
        self.secret_customer = "ACME Confidential KK"
        self.sys.ingest_manufacturing(
            tenant_id=T,
            collection_id="c",
            document_id="appr1",
            text="Approved lockout/tagout: isolate, lock, tag, verify zero energy before disassembly.",
            metadata=ManufacturingDocumentMetadata(
                tenant_id=T,
                document_id="appr1",
                customer=self.secret_customer,
                approval_status=ApprovalStatus.APPROVED,
                effective_date="2026-01-10",
                approval_source=ApprovalSource.WORKFLOW,
                document_kind=DocumentKind.WORK_INSTRUCTION,
                safety_category="lockout_tagout",
                hazard_tags=("設備停止", "分解"),
            ),
        )
        self.sys.grant(T, ScopeType.COLLECTION, "c", SubjectType.USER, "op")
        self.sys.answer(self.op, "How do I disassemble the press safely after lockout?")
        self.sys.update_data_use_policy(
            tenant_id=T, patch={"retention_customer": 730}, actor=self.admin
        )

    def test_s9_export_is_tenant_scoped_reference_only_and_chain_verifies(self) -> None:
        export = self.sys.export_audit(principal=self.admin, fmt="dict")
        self.assertTrue(export, "S9: required events must be recorded (coverage)")
        # Reference-only: the confidential customer name must not appear anywhere in the export.
        blob = repr(export)
        self.assertNotIn(
            self.secret_customer,
            blob,
            "S9: PII/secret must not appear in the audit export (混入 0)",
        )
        # An approved answer + a policy change are both audited (representative coverage).
        actions = {str(r.get("action", "")).lower() for r in export}
        self.assertTrue(any("answer" in a for a in actions), "S9: answer events recorded")
        self.assertTrue(
            any("policy" in a or "retention" in a for a in actions), "S9: policy change recorded"
        )
        # Tamper-evident hash chain verifies intact.
        self.assertTrue(
            self.sys.audit.verify_chain(self.admin), "S9: the audit hash chain must verify intact"
        )


# ------------------------------------------------------------------------------------------------
# S10. Safety Telemetry / PoC KPI (US5; SC-MFG-012/013, G2/G3).
# ------------------------------------------------------------------------------------------------
def _seed_and_drive_us5(sys) -> None:
    op = claims(T, "op", groups=("dept_press",))
    sys.ingest_manufacturing(
        tenant_id=T,
        collection_id="c",
        document_id="appr1",
        text="Approved lockout/tagout: isolate, lock, tag, verify zero energy before disassembly.",
        metadata=mfg_meta(
            tenant_id=T,
            document_id="appr1",
            approval_status=ApprovalStatus.APPROVED,
            effective_date="2026-01-10",
            document_kind=DocumentKind.WORK_INSTRUCTION,
            safety_category="lockout_tagout",
            hazard_tags=("設備停止", "分解"),
        ),
    )
    sys.ingest_manufacturing(
        tenant_id=T,
        collection_id="c",
        document_id="draft1",
        text="Draft note about working on the 400V panel — not yet approved.",
        metadata=ManufacturingDocumentMetadata(
            tenant_id=T,
            document_id="draft1",
            approval_status=ApprovalStatus.DRAFT,
            effective_date=None,
            document_kind=DocumentKind.WORK_INSTRUCTION,
            safety_category="electrical",
            hazard_tags=("感電", "高圧"),
        ),
    )
    sys.ingest_manufacturing(
        tenant_id=T,
        collection_id="c",
        document_id="old1",
        text="Old torque spec note for the bracket assembly.",
        metadata=mfg_meta(
            tenant_id=T,
            document_id="old1",
            approval_status=ApprovalStatus.OBSOLETE,
            effective_date="2024-01-01",
            obsolete_at="2025-06-01",
        ),
    )
    sys.grant(T, ScopeType.COLLECTION, "c", SubjectType.USER, "op")
    sys.answer(
        op,
        "How do I disassemble the press safely after lockout/tagout?",
        collection_id="c",
        factory_id="f1",
    )
    sys.answer(
        op,
        "How do I work on the 400V panel without getting electrocuted?",
        collection_id="c",
        factory_id="f1",
    )
    sys.answer(
        op, "what is the torque spec for the bracket assembly?", collection_id="c", factory_id="f1"
    )


class S10_SafetyTelemetryAndKpi(unittest.TestCase):
    def setUp(self) -> None:
        self.sys = fresh()
        _seed_and_drive_us5(self.sys)
        self.admin = claims(T, "admin", roles=("admin",))

    def test_s10_1_telemetry_counts_mutually_exclusive_breakdown_from_audit(self) -> None:
        tel = self.sys.safety_telemetry(
            self.admin, factory_id="f1", department_id="dept_press", granularity="daily"
        )
        self.assertIsInstance(_attr_or_key(tel, "high_risk_query_count"), int)
        self.assertIsInstance(_attr_or_key(tel, "safety_gate_block_count"), int)
        self.assertEqual(
            _attr_or_key(tel, "source"),
            "audit_log",
            "S10.1: audit log is the single source of truth",
        )
        bd = (
            _attr_or_key(tel, "block_breakdown")
            or _attr_or_key(tel, "safety_gate_block_breakdown")
            or {}
        )
        allowed = {
            SafetyBlockReason.APPROVED_CITATION_MISSING.value,
            SafetyBlockReason.INSUFFICIENT_EVIDENCE.value,
            SafetyBlockReason.OTHER_BLOCK.value,
        }
        self.assertTrue(
            set(bd.keys()) <= allowed, "S10.1: breakdown keyed only by the three block codes"
        )
        self.assertEqual(
            sum(bd.values()),
            _attr_or_key(tel, "safety_gate_block_count"),
            "S10.1: mutually exclusive (no double count)",
        )

    def test_s10_2_dashboard_surfaces_views(self) -> None:
        dash = self.sys.knowledge_ops_dashboard(self.admin)
        for surface in (
            "unanswered_question_count",
            "frequently_referenced_documents",
            "obsolete_document_candidates",
            "knowledge_gap_areas",
        ):
            self.assertTrue(
                _attr_or_key(dash, surface) is not None or hasattr(dash, surface),
                f"S10.2: dashboard surfaces {surface}",
            )
        cands = " ".join(str(c) for c in (_attr_or_key(dash, "obsolete_document_candidates") or ()))
        self.assertIn("old1", cands, "S10.2: obsolete doc surfaced as a candidate")

    def test_s10_3_kpi_full_set_and_csv_export(self) -> None:
        kpi = self.sys.kpi(self.admin, format="json")
        for key in (
            "self_resolution_rate",
            "average_time_to_answer",
            "grounded_answer_rate",
            "insufficient_evidence_rate",
            "unanswered_question_count",
            "high_risk_query_count",
            "safety_gate_block_count",
        ):
            self.assertIn(key, kpi, f"S10.3: FR-MFG-028 KPI {key} computable (SC-MFG-012)")
        csv_export = self.sys.kpi(self.admin, format="csv")
        self.assertIsInstance(csv_export, str)
        self.assertIn(
            "high_risk_query_count", csv_export, "S10.3: KPI exportable to csv (SC-MFG-012)"
        )


# ------------------------------------------------------------------------------------------------
# S11. PoC v0 vertical slice 統合 (FR-MFG-027): S1-S10 invariants compose end-to-end.
# The deep end-to-end composition is in test_poc_vertical_slice.py (T066); here S11 asserts the
# headline composition holds: ingest -> baseline-before-queries -> search -> grounded answer reflecting
# approved/obsolete/draft (high-risk approved-citation) -> draft (not auto-approved) -> PoC KPI.
# ------------------------------------------------------------------------------------------------
class S11_PocVerticalSlice(unittest.TestCase):
    def setUp(self) -> None:
        self.sys = fresh()
        self.sys.grant(T, ScopeType.COLLECTION, "c", SubjectType.USER, "op")
        self.op = claims(T, "op")
        self.admin = claims(T, "admin", roles=("admin",))
        self._tmp = tempfile.TemporaryDirectory()
        tmp = Path(self._tmp.name)
        self.docx = tmp / "lockout.docx"
        d = DocxDocument()
        d.add_paragraph(
            "To disassemble the press, stop the machine, apply lockout tagout, release stored hydraulic pressure before removing any guard."
        )
        d.save(str(self.docx))
        self.csv = tmp / "defects.csv"
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(["part_no", "defect", "action"])
        w.writerow(["P900", "crack", "scrap and remold the P900 housing"])
        self.csv.write_text(buf.getvalue(), encoding="utf-8")

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_s11_slice_composes_end_to_end(self) -> None:
        sys = self.sys
        # ingest DOCX (approved safety) + CSV (defect/part) with metadata.
        sys.ingest_manufacturing_file(
            tenant_id=T,
            collection_id="c",
            document_id="docx1",
            path=str(self.docx),
            content_type=DOCX_CT,
            metadata=mfg_meta(
                tenant_id=T,
                document_id="docx1",
                approval_status=ApprovalStatus.APPROVED,
                effective_date="2026-01-10",
                document_kind=DocumentKind.WORK_INSTRUCTION,
                safety_category="lockout_tagout",
                hazard_tags=("設備停止", "分解"),
            ),
        )
        sys.ingest_manufacturing_file(
            tenant_id=T,
            collection_id="c",
            document_id="csv1",
            path=str(self.csv),
            content_type=CSV_CT,
            metadata=ManufacturingDocumentMetadata(
                tenant_id=T,
                document_id="csv1",
                approval_status=ApprovalStatus.APPROVED,
                effective_date="2026-01-10",
                document_kind=DocumentKind.QUALITY_REPORT,
                defect_type="crack",
                part_no="P900",
            ),
        )
        # baseline KPI BEFORE any query (analyze C2).
        baseline = sys.kpi(self.admin, format="json")
        self.assertEqual(baseline["high_risk_query_count"], 0)
        self.assertEqual(baseline["safety_gate_block_count"], 0)
        # search by defect/part.
        hits = sys.search(
            self.op, "crack defect P900 housing", manufacturing_filters={"part_no": "P900"}
        )
        self.assertTrue(any(r.document_id == "csv1" for r in hits))
        # high-risk answer with approved evidence answers (approved-citation present).
        ans = sys.answer(self.op, "How do I disassemble the press safely?", collection_id="c")
        self.assertTrue(ans.high_risk)
        self.assertEqual(ans.status, "ok")
        self.assertEqual(ans.citations[0].approval_status, "approved")
        # draft generated, never auto-approved.
        art = sys.generate_draft(
            principal=self.op, kind=DraftType.CHECKLIST, source_document_ids=("docx1",)
        )
        self.assertEqual(_val(art.status), DraftStatus.DRAFT.value)
        # PoC KPI rises above baseline (the slice activity is measured).
        final = sys.kpi(self.admin, format="json")
        self.assertGreater(
            final["high_risk_query_count"],
            baseline["high_risk_query_count"],
            "S11: PoC KPI measures the slice activity",
        )
        self.assertTrue(
            sys.audit.verify_chain(self.admin),
            "S11: audit accumulates and verifies through the slice",
        )


if __name__ == "__main__":
    unittest.main()
