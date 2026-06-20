"""T066 — PoC v0 vertical slice END-TO-END integration (FR-MFG-027; quickstart S11).

This is the CAPSTONE composition test. It does NOT reimplement any behaviour: it COMPOSES the whole
stack already built on ``raku_rag.manufacturing.app.ManufacturingSystem`` and asserts the full
FR-MFG-027 slice holds together end-to-end on a fresh system with seeded factory data:

  ingest (DOCX/XLSX/CSV + text) with manufacturing metadata + approval/import
    -> capture a BASELINE PoC-KPI snapshot BEFORE any query (analyze C2: baseline-before-queries)
    -> search by equipment name / alarm code / process name / defect type / part number
    -> grounded answer reflecting approved/obsolete/draft document state (high-risk REQUIRES an
       approved+effective citation, else insufficient_evidence with safety_block_reason=
       approved_citation_missing)
    -> similar past trouble-cases shown as CANDIDATE countermeasures (Hard Rule 4)
    -> a generated DraftArtifact is status=draft and only a reviewer can approve it (never auto-confirmed)
    -> PoC KPI + safety telemetry compute, and the telemetry counts MATCH an independent audit scan.

Audit accumulates throughout; the slice is the union of US1 (safety gate), US2 (mfg ingest), US3
(trouble cases), US4 (drafts), US5 (KPI/telemetry/dashboard), US6 (ACL) + Governance, exercised
through the SINGLE composition root.

The invariant *correctness* of each sub-system is mechanism-pinned by its own dedicated hard-gate
test (test_safety_gate / test_obsolete_draft_evidence / test_draft_only / test_trouble_cases /
test_safety_telemetry / test_audit_coverage). This test pins that the slice COMPOSES: every stage
runs in sequence on one system and the cross-stage invariants (baseline-before-queries, high-risk
approved-citation, draft-not-auto-confirmed, telemetry==audit) hold simultaneously.

stdlib only (+ python-docx/openpyxl for the DOCX/XLSX fixtures, already present and used by
test_ingest_formats.py). Fixtures are generated programmatically in a tmp dir (no committed binaries).

Authoritative: spec FR-MFG-027/028, SC-MFG-012/013; quickstart S11 (and S1-S10 invariants);
contracts/mfg-openapi.md. Assertion style mirrors tests/manufacturing/test_ingest_formats.py +
test_dashboard_contract.py + test_safety_telemetry.py.
"""

from __future__ import annotations

import csv
import io
import tempfile
import unittest
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import Workbook

from raku_rag.domain.models import ScopeType, SubjectType
from raku_rag.manufacturing.domain.draft import CreatedBy, DraftStatus, DraftType
from raku_rag.manufacturing.domain.entities import (
    Countermeasure,
    CountermeasureType,
    FailureMode,
    MeasureClass,
    TroubleCase,
)
from raku_rag.manufacturing.domain.metadata import (
    ApprovalSource,
    ApprovalStatus,
    DocumentKind,
    ManufacturingDocumentMetadata,
)
from raku_rag.manufacturing.domain.safety import SafetyBlockReason
from tests.manufacturing.helpers import T, claims, fresh, mfg_meta

DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_CT = "text/csv"
COLL = "c"


# --- audit-side ground truth (telemetry is checked against an INDEPENDENT scan; mirrors T047) ------
def _audit_high_risk_count(entries) -> int:
    return sum(1 for e in entries if getattr(e, "high_risk_classification_result", None) is True)


def _audit_block_count(entries) -> int:
    return sum(1 for e in entries if getattr(e, "safety_block_reason", None) is not None)


def _attr_or_key(obj, name):
    if isinstance(obj, dict):
        return obj.get(name)
    return getattr(obj, name, None)


# ------------------------------------------------------------------------------------------------
# Programmatic fixtures (no committed binaries) — a DOCX work instruction, an XLSX equipment ledger,
# and a CSV defect log. Each carries vocabulary that the equipment/alarm/process/defect/part search
# and the grounded-answer / high-risk stages key off.
# ------------------------------------------------------------------------------------------------
def write_docx(path: Path) -> None:
    """An APPROVED, safety-tagged lockout/tagout work instruction (high-risk positive-control source)."""
    doc = DocxDocument()
    doc.add_paragraph(
        "To disassemble the press, first stop the machine, apply lockout tagout, and release the "
        "stored hydraulic pressure before removing any guard."
    )
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Step"
    table.cell(0, 1).text = "Detail"
    table.cell(1, 0).text = "verify"
    table.cell(1, 1).text = "verify zero energy on the press before disassembly"
    doc.save(str(path))


def write_xlsx(path: Path) -> None:
    """An equipment ledger: pump17 / alarm E152 remedy (alarm-code + equipment search source)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ledger"
    ws.append(["equipment", "alarm", "remedy"])
    ws.append(["pump17", "E152", "replace the impeller seal on pump17 to clear alarm E152"])
    wb.save(str(path))


def write_csv(path: Path) -> None:
    """A defect log: part P900 crack defect action (defect-type + part-number search source)."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["part_no", "defect", "action"])
    w.writerow(["P900", "crack", "scrap and remold the P900 housing when a crack defect is found"])
    path.write_text(buf.getvalue(), encoding="utf-8")


# The seeded trouble-case body (gearbox vibration + abnormal noise) — drives the US3 stage.
TC_BODY = (
    "Trouble report: the assembly gearbox showed vibration increase and abnormal noise. Root cause: "
    "bearing wear. Provisional countermeasure: reduce feed rate and monitor closely. Permanent "
    "countermeasure: replace the worn bearing and install a vibration sensor. Recurrence prevention: "
    "add the bearing to the periodic replacement schedule."
)
SYMPTOM_QUERY = "振動増加 異音 vibration increase abnormal noise gearbox bearing"


class PocVerticalSliceTest(unittest.TestCase):
    """ONE end-to-end run of the FR-MFG-027 slice; per-stage assertions are recorded as we go."""

    def setUp(self) -> None:
        self.sys = fresh()
        self._tmp = tempfile.TemporaryDirectory()
        tmp = Path(self._tmp.name)
        self.docx = tmp / "lockout.docx"
        self.xlsx = tmp / "ledger.xlsx"
        self.csv = tmp / "defects.csv"
        write_docx(self.docx)
        write_xlsx(self.xlsx)
        write_csv(self.csv)
        # One operator with a collection grant covering the whole PoC collection (US6 reuse).
        self.sys.grant(T, ScopeType.COLLECTION, COLL, SubjectType.USER, "op")
        self.op = claims(T, "op")
        self.admin = claims(T, "admin", roles=("admin",))

    def tearDown(self) -> None:
        self._tmp.cleanup()

    # ------------------------------------------------------------------------------------------
    # THE SLICE: one method runs every stage in order on the SAME system.
    # ------------------------------------------------------------------------------------------
    def test_poc_v0_vertical_slice_end_to_end(self) -> None:
        sys = self.sys

        # === STAGE 1 — INGEST DOCX/XLSX/CSV + text, with manufacturing metadata + approval/import ===
        # 1a) APPROVED + effective DOCX safety work instruction (file ingestion path, US2).
        job_docx = sys.ingest_manufacturing_file(
            tenant_id=T,
            collection_id=COLL,
            document_id="docx_lockout",
            path=str(self.docx),
            content_type=DOCX_CT,
            metadata=mfg_meta(
                tenant_id=T,
                document_id="docx_lockout",
                approval_status=ApprovalStatus.APPROVED,
                effective_date="2026-01-10",
                approval_source=ApprovalSource.IMPORTED,
                document_kind=DocumentKind.WORK_INSTRUCTION,
                safety_category="lockout_tagout",
                hazard_tags=("設備停止", "分解", "高圧"),
                equipment_id="press1",
                process_id="proc_press",
            ),
        )
        self.assertEqual(job_docx.status, "succeeded", getattr(job_docx, "failure_reason", ""))
        self.assertGreater(
            job_docx.chunk_count, 0, "DOCX must chunk/index (US2 reuse of 001 ingest)"
        )

        # 1b) XLSX equipment ledger — equipment 'pump17' + alarm 'E152' (search source).
        job_xlsx = sys.ingest_manufacturing_file(
            tenant_id=T,
            collection_id=COLL,
            document_id="xlsx_ledger",
            path=str(self.xlsx),
            content_type=XLSX_CT,
            metadata=ManufacturingDocumentMetadata(
                tenant_id=T,
                document_id="xlsx_ledger",
                approval_status=ApprovalStatus.APPROVED,
                effective_date="2026-01-10",
                document_kind=DocumentKind.LEDGER,
                equipment="pump17",
                alarm_code="E152",
            ),
        )
        self.assertEqual(job_xlsx.status, "succeeded", getattr(job_xlsx, "failure_reason", ""))
        self.assertGreater(job_xlsx.chunk_count, 0)

        # 1c) CSV defect log — defect 'crack' + part 'P900' (search source).
        job_csv = sys.ingest_manufacturing_file(
            tenant_id=T,
            collection_id=COLL,
            document_id="csv_defects",
            path=str(self.csv),
            content_type=CSV_CT,
            metadata=ManufacturingDocumentMetadata(
                tenant_id=T,
                document_id="csv_defects",
                approval_status=ApprovalStatus.APPROVED,
                effective_date="2026-01-10",
                document_kind=DocumentKind.QUALITY_REPORT,
                defect_type="crack",
                part_no="P900",
            ),
        )
        self.assertEqual(job_csv.status, "succeeded", getattr(job_csv, "failure_reason", ""))
        self.assertGreater(job_csv.chunk_count, 0)

        # 1d) TEXT manuals with DISTINCT document states (approved / obsolete / draft) — used by the
        #     answer stage to prove approved/obsolete/draft handling is reflected (FR-MFG-006).
        sys.ingest_manufacturing(
            tenant_id=T,
            collection_id=COLL,
            document_id="approved_torque",
            text="The torque specification for the M8 conveyor cover bolt is twelve newton meters.",
            metadata=ManufacturingDocumentMetadata(
                tenant_id=T,
                document_id="approved_torque",
                approval_status=ApprovalStatus.APPROVED,
                effective_date="2026-01-10",
                approval_source=ApprovalSource.IMPORTED,
                process="conveyor_assembly",
                process_id="proc_conveyor",
            ),
        )
        # An OBSOLETE doc on a topic with NO approved counterpart -> obsolete-ONLY evidence (S4).
        sys.ingest_manufacturing(
            tenant_id=T,
            collection_id=COLL,
            document_id="obsolete_coolant",
            text="The legacy coolant flow rate setpoint for the grinder spindle is eight liters per minute.",
            metadata=mfg_meta(
                tenant_id=T,
                document_id="obsolete_coolant",
                approval_status=ApprovalStatus.OBSOLETE,
                effective_date="2020-01-01",
                obsolete_at="2025-12-31",
            ),
        )
        sys.ingest_manufacturing(
            tenant_id=T,
            collection_id=COLL,
            document_id="draft_panel",
            text="Draft note about working on the 400V electrical panel — not yet approved.",
            metadata=mfg_meta(
                tenant_id=T,
                document_id="draft_panel",
                approval_status=ApprovalStatus.DRAFT,
                effective_date=None,
                document_kind=DocumentKind.WORK_INSTRUCTION,
                safety_category="electrical",
                hazard_tags=("感電", "高圧"),
            ),
        )
        # 1e) An IMPORTED upstream approval as source of truth (FR-MFG-004a) — exercises the import path.
        sys.ingest_manufacturing(
            tenant_id=T,
            collection_id=COLL,
            document_id="imported_doc",
            text="Imported approved maintenance bulletin for the conveyor drive unit.",
            metadata=mfg_meta(
                tenant_id=T,
                document_id="imported_doc",
                approval_status=ApprovalStatus.PENDING_REVIEW,
                effective_date=None,
            ),
        )
        imported_state = sys.import_external_approval(
            tenant_id=T,
            document_id="imported_doc",
            external={
                "approval_status": "approved",
                "effective_date": "2026-01-05",
                "approved_by": "upstream_dms",
            },
            actor=self.admin,
        )
        self.assertEqual(
            getattr(imported_state, "approval_status", None)
            or _attr_or_key(imported_state, "approval_status"),
            "approved",
            "an imported upstream approval becomes the source of truth (FR-MFG-004a)",
        )

        # === STAGE 2 — BASELINE PoC-KPI SNAPSHOT *BEFORE* ANY QUERY (analyze C2) ===================
        # Capture KPI before driving any answer/search/trouble query; the safety counters must be
        # zero at baseline so a post-flow delta is attributable to the slice activity.
        baseline_kpi = sys.kpi(self.admin, format="json")
        self.assertIsInstance(baseline_kpi, dict)
        self.assertEqual(
            baseline_kpi["high_risk_query_count"],
            0,
            "BASELINE captured before any query MUST show 0 high-risk queries (analyze C2)",
        )
        self.assertEqual(
            baseline_kpi["safety_gate_block_count"],
            0,
            "BASELINE captured before any query MUST show 0 safety blocks (analyze C2)",
        )
        self.assertIn("materialized_at", baseline_kpi, "the baseline snapshot must be timestamped")
        baseline_tel = sys.safety_telemetry(self.admin)
        self.assertEqual(_attr_or_key(baseline_tel, "high_risk_query_count"), 0)
        self.assertEqual(_attr_or_key(baseline_tel, "safety_gate_block_count"), 0)

        # === STAGE 3 — SEARCH by equipment / alarm / process / defect / part (FR-MFG-027) =========
        # equipment name (XLSX ledger) — found via 001 retrieval, carries the approval tag.
        eq_hits = sys.search(
            self.op, "replace impeller seal pump17", manufacturing_filters={"equipment": "pump17"}
        )
        self.assertTrue(
            any(r.document_id == "xlsx_ledger" for r in eq_hits),
            "equipment-name search must find the ledger",
        )
        # alarm code (XLSX ledger).
        alarm_hits = sys.search(
            self.op, "alarm E152 remedy", manufacturing_filters={"alarm_code": "E152"}
        )
        self.assertTrue(
            any(r.document_id == "xlsx_ledger" for r in alarm_hits),
            "alarm-code search must find the ledger",
        )
        # process name (approved torque doc).
        proc_hits = sys.search(
            self.op, "torque specification", manufacturing_filters={"process": "conveyor_assembly"}
        )
        self.assertTrue(
            any(r.document_id == "approved_torque" for r in proc_hits),
            "process-name search must find the doc",
        )
        # defect type (CSV defect log).
        defect_hits = sys.search(
            self.op,
            "crack defect P900 scrap remold",
            manufacturing_filters={"defect_type": "crack"},
        )
        self.assertTrue(
            any(r.document_id == "csv_defects" for r in defect_hits),
            "defect-type search must find the CSV",
        )
        # part number (CSV defect log).
        part_hits = sys.search(
            self.op, "P900 housing action", manufacturing_filters={"part_no": "P900"}
        )
        self.assertTrue(
            any(r.document_id == "csv_defects" for r in part_hits),
            "part-number search must find the CSV",
        )
        # The metadata filter is a candidate post-filter; it MUST NOT leak a non-matching doc.
        for r in eq_hits:
            self.assertEqual(r.approval_status is None or isinstance(r.approval_status, str), True)
        self.assertTrue(
            all(r.document_id != "csv_defects" for r in eq_hits),
            "the equipment filter must restrict results to the matching metadata",
        )

        # === STAGE 4 — GROUNDED ANSWER reflecting approved / obsolete / draft state ================
        # 4a) APPROVED non-high-risk question answers with an APPROVED PRIMARY citation (S2). The
        #     approved doc backs the assertion; no obsolete material is among the torque candidates.
        ans_ok = sys.answer(
            self.op,
            "what is the torque specification for the M8 conveyor cover bolt?",
            collection_id=COLL,
        )
        self.assertEqual(ans_ok.status, "ok")
        self.assertTrue(ans_ok.citations)
        self.assertEqual(
            ans_ok.citations[0].document_id,
            "approved_torque",
            "the primary citation must be the approved doc",
        )
        self.assertEqual(ans_ok.citations[0].approval_status, "approved")
        self.assertIsNone(ans_ok.safety_block_reason)
        self.assertFalse(
            ans_ok.obsolete_warning, "an approved answer with no obsolete candidate must not warn"
        )

        # 4b) OBSOLETE evidence raises the mandatory warning when referenced (S4 / FR-MFG-006). The
        #     grinder-coolant topic's relevant evidence is the obsolete doc; the answer must surface
        #     the obsolete_warning (the obsolete-is-never-PRIMARY axis is pinned, with a discovered
        #     integration gap, by ObsoleteNotPrimaryCitationGapTest below).
        ans_obsolete = sys.answer(
            self.op,
            "what is the legacy coolant flow rate setpoint for the grinder spindle?",
            collection_id=COLL,
        )
        self.assertTrue(
            ans_obsolete.obsolete_warning,
            "referencing obsolete evidence requires obsolete_warning (FR-MFG-006)",
        )

        # 4c) HIGH-RISK with an APPROVED+effective citation answers + requires on-site confirmation (S3.3).
        ans_hr_ok = sys.answer(
            self.op, "How do I disassemble the press safely?", collection_id=COLL
        )
        self.assertTrue(ans_hr_ok.high_risk, "press disassembly is a high-risk query")
        self.assertEqual(
            ans_hr_ok.status,
            "ok",
            "an approved+effective safety citation allows the high-risk answer",
        )
        self.assertIsNone(ans_hr_ok.safety_block_reason)
        self.assertTrue(ans_hr_ok.citations)
        self.assertEqual(ans_hr_ok.citations[0].approval_status, "approved")
        self.assertTrue(
            ans_hr_ok.requires_onsite_confirmation,
            "high-risk hazardous work requires on-site confirmation (FR-MFG-007)",
        )

        # 4d) HIGH-RISK WITHOUT an approved+effective citation is BLOCKED with the exact reason (S3.1 / SC-MFG-006).
        #     The only electrical-panel evidence is the DRAFT doc => approved_citation_missing.
        ans_hr_block = sys.answer(
            self.op,
            "How do I work on the 400V panel without getting electrocuted?",
            collection_id=COLL,
        )
        self.assertTrue(ans_hr_block.high_risk)
        self.assertEqual(ans_hr_block.status, "insufficient_evidence")
        self.assertEqual(
            ans_hr_block.safety_block_reason, SafetyBlockReason.APPROVED_CITATION_MISSING.value
        )
        self.assertFalse(ans_hr_block.text, "must not assert a procedure without approved evidence")
        self.assertEqual(ans_hr_block.used_chunks, ())

        # === STAGE 5 — SIMILAR TROUBLE-CASES as candidate countermeasures (Hard Rule 4) ============
        sys.register_trouble_case(
            tenant_id=T,
            collection_id=COLL,
            source_document_id="tr_gearbox",
            text=TC_BODY,
            metadata=mfg_meta(tenant_id=T, document_id="tr_gearbox"),
            trouble_case=TroubleCase(
                tenant_id=T,
                trouble_case_id="tc_gearbox",
                symptom="振動増加＋異音",
                equipment_id="eq_gearbox",
                process_id="proc_assembly",
                failure_mode_id="fm_bearing",
                source_document_id="tr_gearbox",
            ),
            failure_mode=FailureMode(
                tenant_id=T,
                failure_mode_id="fm_bearing",
                name="軸受摩耗",
                description="bearing wear",
            ),
            countermeasures=(
                Countermeasure(
                    tenant_id=T,
                    measure_id="m_prov",
                    trouble_case_id="tc_gearbox",
                    description="reduce feed rate and monitor closely",
                    measure_class=MeasureClass.PROVISIONAL,
                    source_document_id="tr_gearbox",
                ),
                Countermeasure(
                    tenant_id=T,
                    measure_id="m_perm",
                    trouble_case_id="tc_gearbox",
                    description="replace the worn bearing and install a vibration sensor",
                    measure_class=MeasureClass.PERMANENT,
                    source_document_id="tr_gearbox",
                ),
            ),
            recurrence_prevention="add the bearing to the periodic replacement schedule",
        )
        tc_resp = sys.search_trouble_cases(self.op, SYMPTOM_QUERY)
        self.assertEqual(getattr(tc_resp.status, "value", tc_resp.status), "ok")
        match = next((m for m in tc_resp.results if m.trouble_case_id == "tc_gearbox"), None)
        self.assertIsNotNone(
            match, "the seeded gearbox case must be retrieved for the symptom query"
        )
        self.assertIsNotNone(match.failure_mode, "the cause (FailureMode) must accompany the match")
        self.assertTrue(match.recurrence_prevention)
        self.assertTrue(match.citations, "the similar case must be cited (出典つき)")
        # Hard Rule 4: EVERY past-case countermeasure (even the permanent one) displays as candidate.
        all_cms = tuple(match.countermeasures.provisional) + tuple(match.countermeasures.permanent)
        self.assertTrue(all_cms, "the trouble case must surface countermeasures")
        for c in all_cms:
            self.assertEqual(
                getattr(c.type, "value", c.type),
                CountermeasureType.CANDIDATE.value,
                "Hard Rule 4: a past-case countermeasure must display as candidate, never definitive",
            )
            self.assertTrue(c.label, "every candidate carries a past-example label")

        # === STAGE 6 — DraftArtifact: created draft, reviewer-approved, NEVER auto-confirmed =======
        from raku_rag.domain.models import Citation

        draft = sys.generate_draft(
            principal=self.op,
            kind=DraftType.CHECKLIST,
            context_citations=(
                Citation(
                    kind="text",
                    document_id="docx_lockout",
                    source_id="src",
                    version=1,
                    retrieval_score=0.9,
                    chunk_id="docx_lockout#0",
                ),
            ),
            source_document_ids=("docx_lockout",),
        )
        self.assertEqual(
            getattr(draft.status, "value", draft.status),
            DraftStatus.DRAFT.value,
            "AI output is always draft (Hard Rule 1)",
        )
        self.assertEqual(getattr(draft.created_by, "value", draft.created_by), CreatedBy.AI.value)
        # A no-reviewer approve attempt MUST NOT confirm the draft (SC-MFG-007).
        leaked = False
        try:
            r = sys.review_draft(
                tenant_id=T, artifact_id=draft.artifact_id, reviewer=None, decision="approved"
            )
            leaked = getattr(r.status, "value", r.status) == DraftStatus.APPROVED.value
        except (ValueError, PermissionError, TypeError):
            pass
        self.assertFalse(leaked, "AI self-approve (no reviewer) must be rejected (SC-MFG-007 = 0)")
        self.assertEqual(
            getattr(sys.get_draft(T, draft.artifact_id).status, "value", None),
            DraftStatus.DRAFT.value,
        )
        # A reviewer CAN approve it (positive control — approval is reachable only via a reviewer).
        sys.assign_reviewer(tenant_id=T, artifact_id=draft.artifact_id, reviewer_id="rev_poc")
        approved = sys.review_draft(
            tenant_id=T,
            artifact_id=draft.artifact_id,
            reviewer=claims(T, "rev_poc", roles=("reviewer",)),
            decision="approved",
            comment="ok",
        )
        self.assertEqual(
            getattr(approved.status, "value", approved.status), DraftStatus.APPROVED.value
        )
        self.assertEqual(
            approved.reviewer_id, "rev_poc", "the approval is attributable to the human reviewer"
        )

        # === STAGE 7 — PoC KPI + SAFETY TELEMETRY compute; telemetry == an INDEPENDENT audit scan ==
        final_kpi = sys.kpi(self.admin, format="json")
        self.assertIsInstance(final_kpi, dict)
        for key in (
            "self_resolution_rate",
            "average_time_to_answer",
            "grounded_answer_rate",
            "insufficient_evidence_rate",
            "high_risk_query_count",
            "safety_gate_block_count",
        ):
            self.assertIn(key, final_kpi, f"FR-MFG-028 KPI {key!r} must be computable (SC-MFG-012)")
        # CSV export works (SC-MFG-012 export).
        kpi_csv = sys.kpi(self.admin, format="csv")
        self.assertIsInstance(kpi_csv, str)
        self.assertIn("high_risk_query_count", kpi_csv)

        # Safety telemetry counts MATCH an independent scan of the SAME audit log (single source).
        tel = sys.safety_telemetry(self.admin)
        entries = sys.audit.read_all(self.admin)
        self.assertEqual(
            _attr_or_key(tel, "high_risk_query_count"),
            _audit_high_risk_count(entries),
            "telemetry high_risk_query_count MUST equal an independent audit-log scan (single source)",
        )
        self.assertEqual(
            _attr_or_key(tel, "safety_gate_block_count"),
            _audit_block_count(entries),
            "telemetry safety_gate_block_count MUST equal an independent audit-log scan (single source)",
        )
        self.assertEqual(_attr_or_key(tel, "source"), "audit_log")
        # The KPI safety counters agree with the telemetry (consistency across endpoints).
        self.assertEqual(
            final_kpi["high_risk_query_count"], _attr_or_key(tel, "high_risk_query_count")
        )
        self.assertEqual(
            final_kpi["safety_gate_block_count"], _attr_or_key(tel, "safety_gate_block_count")
        )

        # BASELINE-vs-FINAL delta: the slice drove >=1 high-risk query and >=1 safety block, so the
        # counters strictly increased from the pre-query baseline (analyze C2 closure).
        self.assertGreater(
            final_kpi["high_risk_query_count"],
            baseline_kpi["high_risk_query_count"],
            "the slice must drive high-risk queries above the pre-query baseline (analyze C2)",
        )
        self.assertGreater(
            final_kpi["safety_gate_block_count"],
            baseline_kpi["safety_gate_block_count"],
            "the slice must drive a safety block above the pre-query baseline (analyze C2)",
        )

        # Audit accumulated throughout and the tenant chain is intact (tamper-evident, single source).
        self.assertTrue(entries, "audit must accumulate throughout the slice")
        self.assertTrue(
            sys.audit.verify_chain(self.admin), "the accumulated audit chain must verify intact"
        )


# ================================================================================================
# A few FOCUSED slice tests (each pins one cross-stage invariant in isolation, smaller blast radius).
# ================================================================================================
class BaselineBeforeQueriesTest(unittest.TestCase):
    """analyze C2: a PoC-KPI baseline can be captured BEFORE any query (counters start at zero)."""

    def test_baseline_kpi_is_zero_before_any_query(self) -> None:
        sys = fresh()
        admin = claims(T, "admin", roles=("admin",))
        # Ingest only — no answer/search/trouble query has run yet.
        sys.ingest_manufacturing(
            tenant_id=T,
            collection_id=COLL,
            document_id="d1",
            text="The conveyor motor lubrication interval is every ninety days.",
            metadata=mfg_meta(tenant_id=T, document_id="d1"),
        )
        kpi = sys.kpi(admin, format="json")
        self.assertEqual(kpi["high_risk_query_count"], 0)
        self.assertEqual(kpi["safety_gate_block_count"], 0)
        self.assertIn("materialized_at", kpi)
        tel = sys.safety_telemetry(admin)
        self.assertEqual(_attr_or_key(tel, "high_risk_query_count"), 0)
        self.assertEqual(_attr_or_key(tel, "safety_gate_block_count"), 0)


class HighRiskApprovedCitationRequiredAcrossSliceTest(unittest.TestCase):
    """The high-risk approved-citation rule holds when reached through the full ingest->answer slice."""

    def setUp(self) -> None:
        self.sys = fresh()
        self.sys.grant(T, ScopeType.COLLECTION, COLL, SubjectType.USER, "op")
        self.op = claims(T, "op")

    def test_high_risk_blocked_when_only_draft_evidence(self) -> None:
        self.sys.ingest_manufacturing(
            tenant_id=T,
            collection_id=COLL,
            document_id="draft_only",
            text="Draft note about disassembling the press — not yet approved.",
            metadata=mfg_meta(
                tenant_id=T,
                document_id="draft_only",
                approval_status=ApprovalStatus.DRAFT,
                effective_date=None,
                document_kind=DocumentKind.WORK_INSTRUCTION,
                safety_category="lockout_tagout",
                hazard_tags=("設備停止", "分解"),
            ),
        )
        ans = self.sys.answer(self.op, "How do I disassemble the press safely?", collection_id=COLL)
        self.assertTrue(ans.high_risk)
        self.assertEqual(ans.status, "insufficient_evidence")
        self.assertEqual(ans.safety_block_reason, SafetyBlockReason.APPROVED_CITATION_MISSING.value)

    def test_high_risk_answers_when_approved_effective_evidence(self) -> None:
        self.sys.ingest_manufacturing(
            tenant_id=T,
            collection_id=COLL,
            document_id="approved_lockout",
            text="To disassemble the press, stop the machine, apply lockout tagout, release pressure.",
            metadata=mfg_meta(
                tenant_id=T,
                document_id="approved_lockout",
                approval_status=ApprovalStatus.APPROVED,
                effective_date="2026-01-10",
                document_kind=DocumentKind.WORK_INSTRUCTION,
                safety_category="lockout_tagout",
                hazard_tags=("設備停止", "分解"),
            ),
        )
        ans = self.sys.answer(self.op, "How do I disassemble the press safely?", collection_id=COLL)
        self.assertTrue(ans.high_risk)
        self.assertEqual(ans.status, "ok")
        self.assertEqual(ans.citations[0].approval_status, "approved")
        self.assertTrue(ans.requires_onsite_confirmation)


class TelemetryMatchesAuditAfterSliceTest(unittest.TestCase):
    """After driving the answer flow, telemetry counts equal an independent audit-log scan (US5-2)."""

    def test_telemetry_counts_match_audit_scan(self) -> None:
        sys = fresh()
        sys.grant(T, ScopeType.COLLECTION, COLL, SubjectType.USER, "op")
        op = claims(T, "op")
        admin = claims(T, "admin", roles=("admin",))
        sys.ingest_manufacturing(
            tenant_id=T,
            collection_id=COLL,
            document_id="appr",
            text="Approved lockout/tagout: isolate, lock, tag, verify zero energy before disassembly.",
            metadata=mfg_meta(
                tenant_id=T,
                document_id="appr",
                approval_status=ApprovalStatus.APPROVED,
                effective_date="2026-01-10",
                document_kind=DocumentKind.WORK_INSTRUCTION,
                safety_category="lockout_tagout",
                hazard_tags=("設備停止", "分解"),
            ),
        )
        sys.ingest_manufacturing(
            tenant_id=T,
            collection_id=COLL,
            document_id="dft",
            text="Draft note about the 400V panel — not approved.",
            metadata=mfg_meta(
                tenant_id=T,
                document_id="dft",
                approval_status=ApprovalStatus.DRAFT,
                effective_date=None,
                document_kind=DocumentKind.WORK_INSTRUCTION,
                safety_category="electrical",
                hazard_tags=("感電", "高圧"),
            ),
        )
        sys.answer(op, "How do I disassemble the press safely after lockout?", collection_id=COLL)
        sys.answer(
            op, "How do I work on the 400V panel without getting electrocuted?", collection_id=COLL
        )
        tel = sys.safety_telemetry(admin)
        entries = sys.audit.read_all(admin)
        self.assertEqual(
            _attr_or_key(tel, "high_risk_query_count"), _audit_high_risk_count(entries)
        )
        self.assertEqual(_attr_or_key(tel, "safety_gate_block_count"), _audit_block_count(entries))
        self.assertGreater(_attr_or_key(tel, "high_risk_query_count"), 0)
        self.assertGreater(_attr_or_key(tel, "safety_gate_block_count"), 0)


class ObsoleteNotPrimaryCitationGapTest(unittest.TestCase):
    """INTEGRATION GAP discovered by the T066 slice (FR-MFG-006 / SC-MFG-011 docstring invariant).

    The ManufacturingSafetyGate docstring states: "An OBSOLETE document among the candidates =>
    obsolete_warning=True AND it is never PRIMARY evidence. If the ONLY thing that could back an
    assertion is obsolete/draft, the answer cannot be ok." The dedicated SC-MFG-011 hard gate
    (test_obsolete_draft_evidence.py) verifies this when the obsolete doc is the ONLY candidate.

    The slice reveals the gap when OTHER, unrelated APPROVED documents coexist: for a query whose only
    on-topic evidence is obsolete, retrieval still admits loosely-related approved docs (they share
    generic tokens) as candidates. The gate's ``has_usable_primary`` check sees those approved docs and
    does NOT block, so the reused 001 answer path asserts (status=ok) and cites the obsolete doc as the
    TOP (primary) evidence — violating the "obsolete is never primary" invariant the gate promises.

    This test pins the invariant the gate documents. The T066 slice surfaced the gap; the thin
    integration glue added to ``ManufacturingAnswerService.answer`` (api/answer_ext.py) closes it by
    verifying the FINAL/primary cited document is non-obsolete on a non-high-risk ``ok`` answer (not
    merely that SOME candidate is usable-primary): when the asserted answer's top citation is obsolete
    and no approved+effective document is among the cited evidence, the answer is demoted to
    ``insufficient_evidence`` (reference-only with the mandatory obsolete warning). With that glue the
    invariant holds and this test is GREEN; reverting the glue makes it RED again.
    """

    def setUp(self) -> None:
        self.sys = fresh()
        self.sys.grant(T, ScopeType.COLLECTION, COLL, SubjectType.USER, "op")
        self.op = claims(T, "op")
        self._tmp = tempfile.TemporaryDirectory()
        tmp = Path(self._tmp.name)
        xlsx = tmp / "ledger.xlsx"
        csv_path = tmp / "defects.csv"
        write_xlsx(xlsx)
        write_csv(csv_path)
        # The ONLY on-topic (coolant) evidence is OBSOLETE.
        self.sys.ingest_manufacturing(
            tenant_id=T,
            collection_id=COLL,
            document_id="obsolete_coolant",
            text="The legacy coolant flow rate setpoint for the grinder spindle is eight liters per minute.",
            metadata=mfg_meta(
                tenant_id=T,
                document_id="obsolete_coolant",
                approval_status=ApprovalStatus.OBSOLETE,
                effective_date="2020-01-01",
                obsolete_at="2025-12-31",
            ),
        )
        # Unrelated APPROVED spreadsheet/CSV docs whose cell-anchor normalized text shares enough
        # generic tokens to SURVIVE the 001 groundedness pre-gate as weak candidates — the noise that
        # makes the gate's ``has_usable_primary`` check pass while the obsolete doc remains top-scored.
        self.sys.ingest_manufacturing_file(
            tenant_id=T,
            collection_id=COLL,
            document_id="xlsx_noise",
            path=str(xlsx),
            content_type=XLSX_CT,
            metadata=ManufacturingDocumentMetadata(
                tenant_id=T,
                document_id="xlsx_noise",
                approval_status=ApprovalStatus.APPROVED,
                effective_date="2026-01-10",
                document_kind=DocumentKind.LEDGER,
                equipment="pump17",
                alarm_code="E152",
            ),
        )
        self.sys.ingest_manufacturing_file(
            tenant_id=T,
            collection_id=COLL,
            document_id="csv_noise",
            path=str(csv_path),
            content_type=CSV_CT,
            metadata=ManufacturingDocumentMetadata(
                tenant_id=T,
                document_id="csv_noise",
                approval_status=ApprovalStatus.APPROVED,
                effective_date="2026-01-10",
                document_kind=DocumentKind.QUALITY_REPORT,
                defect_type="crack",
                part_no="P900",
            ),
        )

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_obsolete_is_never_the_primary_cited_evidence(self) -> None:
        ans = self.sys.answer(
            self.op,
            "what is the legacy coolant flow rate setpoint for the grinder spindle?",
            collection_id=COLL,
        )
        # The obsolete material is referenced => the mandatory warning fires (FR-MFG-006).
        self.assertTrue(
            ans.obsolete_warning, "obsolete material in the candidates must raise the warning"
        )
        # INVARIANT (FR-MFG-006/SC-MFG-011): an obsolete document is NEVER the primary cited evidence
        # behind an asserted answer. Either the answer does not assert, or its primary citation is not
        # obsolete. (Pre-glue this FAILED: status==ok with obsolete_coolant as citations[0].)
        if ans.status == "ok":
            self.assertNotEqual(
                ans.citations[0].approval_status,
                "obsolete",
                "an obsolete document must never be the PRIMARY citation behind an asserted answer "
                "(FR-MFG-006/SC-MFG-011)",
            )
        else:
            # Demoted to insufficient_evidence: the obsolete doc is reference-only, never the basis.
            self.assertEqual(ans.status, "insufficient_evidence")
            self.assertEqual(ans.safety_block_reason, SafetyBlockReason.INSUFFICIENT_EVIDENCE.value)


if __name__ == "__main__":
    unittest.main()
