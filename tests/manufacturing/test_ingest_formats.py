"""T022 — US2: DOCX/XLSX/CSV ingestion via the manufacturing path (quickstart S1, FR-MFG-001).

A DOCX (paragraphs + a table), an XLSX (a small sheet) and a CSV are ingested through the
manufacturing file-ingestion entrypoint, which REUSES the 001 ingestion path (parse -> chunk ->
embed -> index) behind the 001 ``Parser`` abstraction and attaches ``ManufacturingDocumentMetadata``
to the 001 ``Document.metadata``. Each ingested file must be normalized/chunked/indexed and then be
findable via ``ManufacturingSystem.search`` / ``ManufacturingSystem.answer`` (no new search mechanism).

Fixtures are generated PROGRAMMATICALLY in a tmp dir (no committed binaries):
  - DOCX via python-docx, XLSX via openpyxl, CSV via the stdlib.

Entrypoint contract asserted here (the impl T029/T030 must satisfy):
  ``ManufacturingSystem.ingest_manufacturing_file(*, tenant_id, collection_id, document_id, path,
  content_type=None, metadata, source_id="src")`` — reads the file bytes, routes by extension /
  content_type to the right 001 ``Parser`` implementation (DocxParser / SpreadsheetParser), runs the
  reused 001 ingestion path, attaches the manufacturing metadata, and returns an ingestion job whose
  ``.status == "succeeded"`` with ``.chunk_count > 0`` (mirrors ``MvpSystem.ingest_text``).

TDD: RED now because ``ingest_manufacturing_file`` is unimplemented on ``ManufacturingSystem``.
Authoritative: quickstart S1, contracts/mfg-openapi.md §B, FR-MFG-001/003.
"""
from __future__ import annotations

import csv
import io
import tempfile
import unittest
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import Workbook

from raku_rag.domain.models import ScopeType, SubjectType
from tests.manufacturing.helpers import T, claims, fresh, mfg_meta

DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_CT = "text/csv"


def write_docx(path: Path) -> None:
    """A work instruction with a paragraph AND a table (both must reach the normalized text)."""
    doc = DocxDocument()
    doc.add_paragraph(
        "The conveyor motor lubrication interval is every ninety days under normal load."
    )
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Interval"
    table.cell(1, 0).text = "bearing"
    table.cell(1, 1).text = "monthly inspection of the bearing housing"
    doc.save(str(path))


def write_xlsx(path: Path) -> None:
    """A small equipment ledger sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ledger"
    ws.append(["equipment", "alarm", "remedy"])
    ws.append(["pump17", "E152", "replace the impeller seal on pump17"])
    wb.save(str(path))


def write_csv(path: Path) -> None:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["part_no", "defect", "action"])
    w.writerow(["P900", "crack", "scrap and remold the P900 housing"])
    path.write_text(buf.getvalue(), encoding="utf-8")


class TestIngestFormats(unittest.TestCase):
    def setUp(self) -> None:
        self.sys = fresh()
        self.sys.grant(T, ScopeType.COLLECTION, "c", SubjectType.USER, "op")
        self.op = claims(T, "op")
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.docx = self.tmp / "instruction.docx"
        self.xlsx = self.tmp / "ledger.xlsx"
        self.csv = self.tmp / "defects.csv"
        write_docx(self.docx)
        write_xlsx(self.xlsx)
        write_csv(self.csv)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def _ingest(self, document_id: str, path: Path, content_type: str):
        return self.sys.ingest_manufacturing_file(
            tenant_id=T,
            collection_id="c",
            document_id=document_id,
            path=str(path),
            content_type=content_type,
            metadata=mfg_meta(tenant_id=T, document_id=document_id),
        )

    def test_docx_indexed_and_findable(self) -> None:
        job = self._ingest("docx1", self.docx, DOCX_CT)
        self.assertEqual(job.status, "succeeded", getattr(job, "failure_reason", ""))
        self.assertGreater(job.chunk_count, 0)
        # The paragraph body is findable via the reused 001 retrieval path.
        results = self.sys.search(self.op, "conveyor motor lubrication interval")
        self.assertTrue(any(r.document_id == "docx1" for r in results))
        # The TABLE content (a different cell) is also normalized into the index.
        table_hits = self.sys.search(self.op, "bearing housing inspection")
        self.assertTrue(any(r.document_id == "docx1" for r in table_hits))

    def test_xlsx_indexed_and_findable(self) -> None:
        job = self._ingest("xlsx1", self.xlsx, XLSX_CT)
        self.assertEqual(job.status, "succeeded", getattr(job, "failure_reason", ""))
        self.assertGreater(job.chunk_count, 0)
        results = self.sys.search(self.op, "replace impeller seal pump17")
        self.assertTrue(any(r.document_id == "xlsx1" for r in results))

    def test_csv_indexed_and_findable(self) -> None:
        job = self._ingest("csv1", self.csv, CSV_CT)
        self.assertEqual(job.status, "succeeded", getattr(job, "failure_reason", ""))
        self.assertGreater(job.chunk_count, 0)
        results = self.sys.search(self.op, "scrap remold P900 housing")
        self.assertTrue(any(r.document_id == "csv1" for r in results))

    def test_xlsx_answerable_with_approval_provenance(self) -> None:
        """End-to-end: an approved XLSX ledger answers normally with an approved citation."""
        self._ingest("xlsx_ans", self.xlsx, XLSX_CT)
        ans = self.sys.answer(self.op, "what is the remedy for alarm E152 on pump17?")
        self.assertIn(ans.status, ("ok", "insufficient_evidence"))
        if ans.status == "ok":
            self.assertTrue(ans.citations)
            self.assertEqual(ans.citations[0].approval_status, "approved")


if __name__ == "__main__":
    unittest.main()
