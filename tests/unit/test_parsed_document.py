"""ADR-018 §7/§9 — ParsedDocument contract + structured-parser normalization (Phase B1/B2).

The load-bearing invariant is *parity*: ``StructuredParser.parse_structured(...).text_for_embedding()``
must equal the legacy ``Parser.parse(...)`` string for the same bytes, so moving ingestion onto the
canonical contract cannot silently change chunking / retrieval / citations.
"""

from __future__ import annotations

import importlib.util
import io
import unittest

from raku_rag.domain.parsed_document import (
    ANCHOR_SPREADSHEET_CELL,
    BLOCK_KINDS,
    BLOCK_PARAGRAPH,
    BLOCK_TABLE_CELL,
    PARSED_DOCUMENT_SCHEMA_VERSION,
    Block,
    ParsedDocument,
    Provenance,
    SourceAnchor,
)
from raku_rag.providers.parsers import (
    DOCX_CONTENT_TYPE,
    XLSX_CONTENT_TYPE,
    DocxParser,
    SpreadsheetParser,
    TextParser,
)
from raku_rag.providers.structured_parsers import (
    DocxStructuredParser,
    SpreadsheetStructuredParser,
    TextStructuredParser,
)

_HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None
_HAS_DOCX = importlib.util.find_spec("docx") is not None


class ParsedDocumentContractTest(unittest.TestCase):
    def test_default_schema_version(self) -> None:
        self.assertEqual(ParsedDocument().schema_version, PARSED_DOCUMENT_SCHEMA_VERSION)

    def test_ordered_blocks_respects_reading_order(self) -> None:
        doc = ParsedDocument(
            blocks=(
                Block(block_id="b2", kind=BLOCK_PARAGRAPH, text="second", reading_order=1),
                Block(block_id="b1", kind=BLOCK_PARAGRAPH, text="first", reading_order=0),
            )
        )
        self.assertEqual([b.block_id for b in doc.ordered_blocks()], ["b1", "b2"])

    def test_text_for_embedding_joins_in_reading_order_skipping_empty(self) -> None:
        doc = ParsedDocument(
            blocks=(
                Block(block_id="b0", kind=BLOCK_PARAGRAPH, text="alpha", reading_order=0),
                Block(block_id="b1", kind=BLOCK_PARAGRAPH, text="   ", reading_order=1),
                Block(block_id="b2", kind=BLOCK_PARAGRAPH, normalized_text="beta", reading_order=2),
            )
        )
        self.assertEqual(doc.text_for_embedding(), "alpha\n\nbeta")

    def test_to_dict_is_serializable_and_roundtrips_values(self) -> None:
        doc = ParsedDocument(
            blocks=(
                Block(
                    block_id="b0",
                    kind=BLOCK_TABLE_CELL,
                    text="x",
                    source_anchor=SourceAnchor(
                        type=ANCHOR_SPREADSHEET_CELL, sheet="s", row=2, col=1
                    ),
                    provenance=Provenance(provider="p"),
                ),
            )
        )
        d = doc.to_dict()
        self.assertEqual(d["schema_version"], PARSED_DOCUMENT_SCHEMA_VERSION)
        self.assertEqual(d["blocks"][0]["source_anchor"]["sheet"], "s")
        self.assertEqual(d["blocks"][0]["provenance"]["provider"], "p")

    def test_block_kind_constants_are_registered(self) -> None:
        self.assertIn(BLOCK_PARAGRAPH, BLOCK_KINDS)
        self.assertIn(BLOCK_TABLE_CELL, BLOCK_KINDS)


class TextParityTest(unittest.TestCase):
    def _assert_parity(self, raw: bytes, content_type: str) -> ParsedDocument:
        legacy = TextParser().parse(raw, content_type)
        structured = TextStructuredParser().parse_structured(raw, content_type)
        self.assertEqual(structured.text_for_embedding(), legacy)
        return structured

    def test_plain_text_parity(self) -> None:
        doc = self._assert_parity(b"First paragraph.\n\nSecond paragraph.", "text/plain")
        self.assertTrue(all(b.kind == BLOCK_PARAGRAPH for b in doc.blocks))
        self.assertEqual(len(doc.blocks), 2)

    def test_markdown_collapses_blank_runs_with_parity(self) -> None:
        self._assert_parity(b"# Heading\n\nBody.\n\n\n\nTail.", "text/markdown")

    def test_html_tag_strip_parity(self) -> None:
        self._assert_parity(b"<p>Hello</p><p>World</p>", "text/html")


class SpreadsheetParityTest(unittest.TestCase):
    CSV = b"item,remedy\nP-12,replace valve\n,\nP-13,tighten bolt"

    def test_csv_parity_and_anchors(self) -> None:
        legacy = SpreadsheetParser().parse(self.CSV, "text/csv")
        doc = SpreadsheetStructuredParser().parse_structured(self.CSV, "text/csv")
        self.assertEqual(doc.text_for_embedding(), legacy)
        # every emitted block is a spreadsheet-cell anchor (empty cell skipped => no R3 blocks)
        self.assertTrue(all(b.kind == BLOCK_TABLE_CELL for b in doc.blocks))
        first = doc.blocks[0]
        self.assertEqual(first.source_anchor.type, ANCHOR_SPREADSHEET_CELL)
        self.assertEqual((first.source_anchor.sheet, first.source_anchor.row, first.source_anchor.col),
                         ("sheet1", 2, 1))
        self.assertIn("sheet1!R2C1", first.text)
        # a structured Table is also produced with a header row
        self.assertEqual(len(doc.tables), 1)
        self.assertTrue(any(c.is_header for c in doc.tables[0].cells))

    @unittest.skipUnless(_HAS_OPENPYXL, "openpyxl not installed")
    def test_xlsx_parity(self) -> None:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["検査項目", "結果"])
        ws.append(["外観", "OK"])
        ws.append(["寸法", "NG"])
        buf = io.BytesIO()
        wb.save(buf)
        raw = buf.getvalue()

        legacy = SpreadsheetParser().parse(raw, XLSX_CONTENT_TYPE)
        doc = SpreadsheetStructuredParser().parse_structured(raw, XLSX_CONTENT_TYPE)
        self.assertEqual(doc.text_for_embedding(), legacy)


class DocxParityTest(unittest.TestCase):
    @unittest.skipUnless(_HAS_DOCX, "python-docx not installed")
    def test_docx_parity(self) -> None:
        from docx import Document as _Doc

        d = _Doc()
        d.add_paragraph("作業前に主電源を停止する。")
        table = d.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "項目"
        table.rows[0].cells[1].text = "値"
        table.rows[1].cells[0].text = "電圧"
        table.rows[1].cells[1].text = "200V"
        buf = io.BytesIO()
        d.save(buf)
        raw = buf.getvalue()

        legacy = DocxParser().parse(raw, DOCX_CONTENT_TYPE)
        doc = DocxStructuredParser().parse_structured(raw, DOCX_CONTENT_TYPE)
        self.assertEqual(doc.text_for_embedding(), legacy)


if __name__ == "__main__":
    unittest.main()
